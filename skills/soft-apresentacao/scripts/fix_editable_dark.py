#!/usr/bin/env python3
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor


def rgb(value: str) -> RGBColor:
    cleaned = value.strip().lstrip("#")
    if len(cleaned) != 6:
        raise argparse.ArgumentTypeError("Use uma cor hexadecimal de 6 dígitos, como 000000.")
    try:
        return RGBColor.from_string(cleaned.upper())
    except ValueError as error:
        raise argparse.ArgumentTypeError(str(error)) from error


parser = argparse.ArgumentParser(description="Restaura fundo escuro no PPTX editável exportado do HTML.")
parser.add_argument("input", type=Path)
parser.add_argument("output", type=Path)
parser.add_argument("--background", type=rgb, default=RGBColor(0, 0, 0))
parser.add_argument("--replace-fill", type=rgb, default=RGBColor(255, 255, 255))
args = parser.parse_args()

presentation = Presentation(args.input)
replaced = 0

for slide in presentation.slides:
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = args.background

    for shape in slide.shapes:
        try:
            if shape.fill.type is not None and shape.fill.fore_color.rgb == args.replace_fill:
                shape.fill.solid()
                shape.fill.fore_color.rgb = args.background
                replaced += 1
        except (AttributeError, TypeError):
            continue

args.output.parent.mkdir(parents=True, exist_ok=True)
presentation.save(args.output)
print(f"{len(presentation.slides)} slides, {replaced} preenchimentos corrigidos -> {args.output}")
