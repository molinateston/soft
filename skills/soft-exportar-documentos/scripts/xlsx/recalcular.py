#!/usr/bin/env python3
"""Recalcula as fórmulas de uma planilha pelo LibreOffice e procura erros.

Uso:
    python3 recalcular.py ARQUIVO.xlsx [--limite 30] [--json] [--mostrar 10] [--forcar]

O que faz:
  1. abre a planilha no LibreOffice sem janela, com perfil isolado e
     configurado para recalcular tudo ao abrir, e salva de novo em .xlsx
     mantendo as fórmulas e gravando os valores calculados;
  2. só troca o arquivo original depois que o recálculo terminou e o
     resultado abriu no openpyxl; se algo falhar, o original fica intacto;
  3. abre o resultado com openpyxl em modo valor (data_only=True) e varre
     todas as abas atrás de erro de fórmula: #REF!, #DIV/0!, #VALUE!, #NAME?,
     #N/A, #NUM!, #NULL! (e os códigos Err:NNN do LibreOffice).

Saída: por aba, a contagem de cada tipo de erro e as primeiras células de
cada tipo (--mostrar, padrão 10). Com --json, o mesmo em JSON.

Planilha com vínculo a outro arquivo (xl/externalLinks) perde os vínculos no
recálculo: o LibreOffice não acha o arquivo de fora e troca os valores por
erro. Nesse caso o script recusa, a menos que se passe --forcar.

Recálculo sem erro prova que a fórmula roda, não que ela está certa.

Códigos de saída: 0 sem erro, 1 há erro de fórmula (ou vínculo externo sem
--forcar), 2 argumento errado, 3 falta arquivo, LibreOffice ou openpyxl,
4 o LibreOffice falhou ou passou do limite de tempo.
"""
import argparse
import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ERROS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
ERRO_LO = re.compile(r"^Err:\d{3}$")

FILTROS = {".xlsx": "xlsx:Calc MS Excel 2007 XML", ".xlsm": "xlsm:Calc MS Excel 2007 VBA XML"}

# Perfil do LibreOffice: recalcular sempre ao abrir (0 = sempre, 1 = nunca, 2 = perguntar).
CONFIG_PERFIL = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" \
xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">\
<value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse">\
<value>0</value></prop></item>
</oor:items>
"""


class Falha(Exception):
    def __init__(self, msg, codigo):
        super().__init__(msg)
        self.codigo = codigo


def rodar(cmd, limite, env):
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                            env=env, start_new_session=True)
    try:
        saida, erro = proc.communicate(timeout=limite)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(proc.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        proc.communicate()
        return None, "", ""
    return proc.returncode, saida.decode("utf-8", "replace"), erro.decode("utf-8", "replace")


def recalcular(origem, limite):
    """Recalcula numa pasta temporária e só então troca o original. Devolve nada; levanta Falha."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise Falha("Falta o LibreOffice (comando soffice). Instale o pacote libreoffice-calc.", 3)
    ext = origem.suffix.lower()
    with tempfile.TemporaryDirectory(prefix="recalcular-") as tmp:
        tmp = Path(tmp)
        perfil = tmp / "perfil"
        (perfil / "user").mkdir(parents=True)
        (perfil / "user" / "registrymodifications.xcu").write_text(CONFIG_PERFIL, encoding="utf-8")
        entrada = tmp / "entrada" / ("planilha" + ext)
        entrada.parent.mkdir()
        shutil.copyfile(origem, entrada)
        saida_dir = tmp / "saida"
        saida_dir.mkdir()
        env = dict(os.environ, HOME=str(tmp), SAL_USE_VCLPLUGIN="svp")
        cmd = [soffice, f"-env:UserInstallation={perfil.as_uri()}", "--headless", "--norestore",
               "--nologo", "--nodefault", "--nolockcheck",
               "--convert-to", FILTROS[ext], "--outdir", str(saida_dir), str(entrada)]
        cod, _, err = rodar(cmd, limite, env)
        resultado = saida_dir / ("planilha" + ext)
        if cod is None:
            raise Falha(f"O LibreOffice passou do limite de {limite} s. O original não foi mexido. "
                        "Tente de novo com --limite maior.", 4)
        if cod != 0 or not resultado.is_file() or resultado.stat().st_size == 0:
            raise Falha(f"O LibreOffice não conseguiu recalcular (código {cod}). O original não foi mexido. "
                        f"{err.strip()[:300]}", 4)
        try:
            import openpyxl
            openpyxl.load_workbook(resultado, data_only=True).close()
        except Exception as e:  # noqa: BLE001
            raise Falha(f"O arquivo recalculado não abre no openpyxl ({e}). O original não foi mexido.", 4)
        provisorio = origem.with_name(f".{origem.name}.recalculando")
        shutil.copyfile(resultado, provisorio)
        try:
            shutil.copymode(origem, provisorio)
        except OSError:
            pass
        os.replace(provisorio, origem)


def varrer(caminho, mostrar):
    import openpyxl
    wb_f = openpyxl.load_workbook(caminho, data_only=False)
    formulas = 0
    for ws in wb_f.worksheets:
        for linha in ws.iter_rows():
            for c in linha:
                if c.data_type == "f":
                    formulas += 1
    wb_f.close()

    wb = openpyxl.load_workbook(caminho, data_only=True)
    abas = []
    for ws in wb.worksheets:
        por_tipo = {}
        for linha in ws.iter_rows():
            for c in linha:
                v = c.value
                if isinstance(v, str) and (v in ERROS or ERRO_LO.match(v)):
                    por_tipo.setdefault(v, []).append(c.coordinate)
        total = sum(len(v) for v in por_tipo.values())
        abas.append({
            "aba": ws.title,
            "total": total,
            "erros": {tipo: {"quantidade": len(cels), "celulas": cels[:mostrar],
                             "fora_da_lista": max(0, len(cels) - mostrar)}
                      for tipo, cels in sorted(por_tipo.items())},
        })
    wb.close()
    return formulas, abas


def main():
    ap = argparse.ArgumentParser(description="Recalcula as fórmulas pelo LibreOffice e lista os erros por aba.")
    ap.add_argument("arquivo", help="planilha .xlsx ou .xlsm (é regravada no lugar)")
    ap.add_argument("--limite", type=int, default=30, help="segundos máximos para o LibreOffice (padrão 30)")
    ap.add_argument("--json", action="store_true", help="imprime o resultado em JSON")
    ap.add_argument("--mostrar", type=int, default=10, help="quantas células listar por tipo de erro (padrão 10)")
    ap.add_argument("--forcar", action="store_true",
                    help="recalcula mesmo com vínculo a outro arquivo, aceitando perder os vínculos")
    a = ap.parse_args()

    origem = Path(a.arquivo)
    resultado = {"arquivo": str(origem), "recalculado": False}

    def falhar(msg, codigo):
        if a.json:
            resultado["erro"] = msg
            print(json.dumps(resultado, ensure_ascii=False, indent=2))
        else:
            print(msg, file=sys.stderr)
        return codigo

    if a.limite < 1:
        return falhar("--limite precisa ser um número de segundos positivo.", 2)
    if a.mostrar < 0:
        return falhar("--mostrar não pode ser negativo.", 2)
    if not origem.is_file():
        return falhar(f"Arquivo não encontrado: {origem}", 3)
    if origem.suffix.lower() not in FILTROS:
        return falhar("Formato não aceito: use .xlsx ou .xlsm.", 2)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return falhar("Falta a biblioteca openpyxl. Instale com: pip install openpyxl", 3)
    try:
        with zipfile.ZipFile(origem) as z:
            externos = [n for n in z.namelist() if n.startswith("xl/externalLinks/") and n.endswith(".xml")]
    except zipfile.BadZipFile:
        return falhar(f"{origem} não é uma planilha .xlsx válida (o zip não abre).", 2)
    if externos and not a.forcar:
        return falhar(f"A planilha tem {len(externos)} vínculo(s) com outro arquivo (xl/externalLinks). "
                      "Recalcular troca esses valores por erro e apaga os vínculos. Copie os valores para "
                      "fora antes, ou rode com --forcar aceitando a perda. Nada foi mexido.", 1)

    try:
        recalcular(origem, a.limite)
    except Falha as f:
        return falhar(str(f), f.codigo)
    resultado["recalculado"] = True
    if externos:
        resultado["aviso"] = "vínculos externos podem ter sido perdidos (--forcar)"

    try:
        formulas, abas = varrer(origem, a.mostrar)
    except Exception as e:  # noqa: BLE001
        return falhar(f"Recalculou, mas o openpyxl não conseguiu ler o resultado: {e}", 4)
    total = sum(x["total"] for x in abas)
    resultado.update({"situacao": "com erros" if total else "sem erros", "formulas": formulas,
                      "total_erros": total, "abas": abas})

    if a.json:
        print(json.dumps(resultado, ensure_ascii=False, indent=2))
    else:
        print(f"Recalculado: {origem} ({formulas} fórmulas)")
        if externos:
            print("Aviso: vínculos com outro arquivo podem ter sido perdidos (--forcar).")
        for aba in abas:
            if not aba["total"]:
                print(f'Aba "{aba["aba"]}": sem erros')
                continue
            print(f'Aba "{aba["aba"]}": {aba["total"]} erro(s)')
            for tipo, info in aba["erros"].items():
                cels = ", ".join(info["celulas"])
                resto = f" e mais {info['fora_da_lista']}" if info["fora_da_lista"] else ""
                print(f"  {tipo}: {info['quantidade']}  {cels}{resto}")
        com_erro = sum(1 for x in abas if x["total"])
        if total:
            print(f"Total: {total} erro(s) em {com_erro} aba(s). Conserte e rode de novo.")
        else:
            print("Total: nenhum erro de fórmula. Isso prova que as fórmulas rodam, não que estão certas.")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
