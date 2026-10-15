"""Testes dos scripts de docx do Soft Office.

Rodar de dentro de novo/:
    python3 -m unittest tests.test_docx -v

Gera os documentos de teste na hora (python-docx e XML escrito à mão).
O teste de abertura no LibreOffice é pulado se o soffice não existir.
"""

import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

import docx
from lxml import etree

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(AQUI), "scripts", "docx")
sys.path.insert(0, SCRIPTS)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {
    "w": W,
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "w15": "http://schemas.microsoft.com/office/word/2012/wordml",
    "w16cid": "http://schemas.microsoft.com/office/word/2016/wordml/cid",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
CABECA = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n')
DECL = ('xmlns:w="%s" xmlns:r="%s" xmlns:w14="%s" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" mc:Ignorable="w14"'
        % (W, NS["r"], NS["w14"]))
REL_BASE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
CT_BASE = "application/vnd.openxmlformats-officedocument.wordprocessingml."
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


def w(nome):
    return "{%s}%s" % (W, nome)


def montar(caminho, corpo, extras=(), sect_extra=""):
    """Gera um .docx a partir do pacote padrão do python-docx com o corpo dado.

    extras: lista de (parte, xml, tipo_conteudo, tipo_relacao, rId).
    """
    base = io.BytesIO()
    docx.Document().save(base)
    with zipfile.ZipFile(base) as z:
        arquivos = {n: z.read(n) for n in z.namelist()}
    arquivos["word/document.xml"] = (
        CABECA + "<w:document %s><w:body>%s<w:sectPr>%s"
        '<w:pgSz w:w="11906" w:h="16838"/>'
        '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" '
        'w:header="708" w:footer="708" w:gutter="0"/></w:sectPr></w:body></w:document>'
        % (DECL, corpo, sect_extra)).encode("utf-8")
    if extras:
        ct = etree.fromstring(arquivos["[Content_Types].xml"])
        rels = etree.fromstring(arquivos["word/_rels/document.xml.rels"])
        for parte, xml, tipo, rel, rid in extras:
            arquivos[parte] = (CABECA + xml).encode("utf-8")
            o = etree.SubElement(ct, "{%s}Override" % NS["ct"])
            o.set("PartName", "/" + parte)
            o.set("ContentType", tipo)
            r = etree.SubElement(rels, "{%s}Relationship" % NS["pr"])
            r.set("Id", rid)
            r.set("Type", rel)
            r.set("Target", parte.split("/", 1)[1])
        arquivos["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8")
        arquivos["word/_rels/document.xml.rels"] = etree.tostring(
            rels, xml_declaration=True, encoding="UTF-8")
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", arquivos.pop("[Content_Types].xml"))
        for nome, dados in arquivos.items():
            z.writestr(nome, dados)
    return caminho


def parte(caminho, nome="word/document.xml"):
    with zipfile.ZipFile(caminho) as z:
        return etree.fromstring(z.read(nome))


def bytes_da_parte(caminho, nome):
    with zipfile.ZipFile(caminho) as z:
        return z.read(nome)


def nomes(caminho):
    with zipfile.ZipFile(caminho) as z:
        return z.namelist()


def textos_dos_paragrafos(raiz, com_excluido=False):
    saida = []
    for p in raiz.iter(w("p")):
        pedacos = []
        for el in p.iter(w("t"), w("delText")):
            dono = next(el.iterancestors(w("p")))
            if dono is not p:
                continue
            if el.tag == w("delText") and not com_excluido:
                continue
            pedacos.append(el.text or "")
        saida.append("".join(pedacos))
    return saida


def assinatura_rpr(run):
    rpr = run.find(w("rPr"))
    if rpr is None:
        return ()
    itens = []
    for c in rpr:
        attrs = tuple(sorted((k, v) for k, v in c.attrib.items() if "rsid" not in k))
        itens.append((c.tag, attrs))
    return tuple(sorted(itens))


def formato_por_caractere(raiz):
    """Para cada caractere de texto: (caractere, formatação, cadeia de contêineres até o parágrafo)."""
    saida = []
    for el in raiz.iter(w("t"), w("delText")):
        run = el.getparent()
        cadeia = []
        for a in run.iterancestors():
            if a.tag == w("p"):
                break
            cadeia.append(a.tag)
        fmt = assinatura_rpr(run)
        for ch in el.text or "":
            saida.append((ch, fmt, tuple(cadeia)))
    return saida


def rodar(script, *args, cwd=None):
    return subprocess.run([sys.executable, os.path.join(SCRIPTS, script)] + list(args),
                          capture_output=True, text=True, cwd=cwd, timeout=120)


def r(texto, rpr="", rsid="00A1", preservar=True):
    esp = ' xml:space="preserve"' if preservar else ""
    rp = "<w:rPr>%s</w:rPr>" % rpr if rpr else ""
    return '<w:r w:rsidR="%s">%s<w:t%s>%s</w:t></w:r>' % (rsid, rp, esp, texto)


def contar(raiz, nome):
    return sum(1 for _ in raiz.iter(w(nome)))


def cabecalho_xml(conteudo, raiz="hdr"):
    return "<w:%s %s>%s</w:%s>" % (raiz, DECL, conteudo, raiz)


class Base(unittest.TestCase):
    def setUp(self):
        self.pasta = tempfile.mkdtemp(prefix="teste_docx_")

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def arq(self, nome):
        return os.path.join(self.pasta, nome)

    def abre_no_python_docx(self, caminho):
        return docx.Document(caminho)

    def abre_no_libreoffice(self, caminho):
        if not SOFFICE:
            self.skipTest("soffice não instalado")
        perfil = tempfile.mkdtemp(prefix="perfil_lo_", dir=self.pasta)
        saida = os.path.join(self.pasta, "pdf")
        os.makedirs(saida, exist_ok=True)
        proc = subprocess.run(
            [SOFFICE, "-env:UserInstallation=file://" + perfil, "--headless",
             "--convert-to", "pdf", "--outdir", saida, caminho],
            capture_output=True, text=True, timeout=180)
        pdf = os.path.join(saida, os.path.splitext(os.path.basename(caminho))[0] + ".pdf")
        self.assertTrue(os.path.exists(pdf) and os.path.getsize(pdf) > 0,
                        "LibreOffice não converteu: %s %s" % (proc.stdout, proc.stderr))
        if shutil.which("pdftotext"):
            return subprocess.run(["pdftotext", pdf, "-"], capture_output=True, text=True).stdout
        return ""


# ====================================================================== juntar

class TestJuntarTrechos(Base):
    def juntar(self, corpo, extras=(), sect_extra=""):
        ent = montar(self.arq("ent.docx"), corpo, extras, sect_extra)
        sai = self.arq("sai.docx")
        proc = rodar("juntar_trechos.py", ent, "-o", sai)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        return ent, sai, proc.stdout

    def assert_mesmo_texto_e_formato(self, ent, sai, parte_nome="word/document.xml"):
        a, b = parte(ent, parte_nome), parte(sai, parte_nome)
        self.assertEqual(textos_dos_paragrafos(a, True), textos_dos_paragrafos(b, True))
        self.assertEqual(formato_por_caractere(a), formato_por_caractere(b))

    def test_junta_mesma_formatacao_ignorando_rsid_e_ortografia(self):
        corpo = ("<w:p>" + r("Olá ", "<w:b/>", "0001") + r("mun", "<w:b/>", "0002")
                 + '<w:proofErr w:type="spellStart"/>' + r("do", "<w:b/>", "0003")
                 + '<w:proofErr w:type="spellEnd"/>' + r(" cruel", "<w:b/>", "0004") + "</w:p>")
        ent, sai, saida = self.juntar(corpo)
        self.assertIn("Juntei 3 trechos", saida)
        self.assertIn("Removi 2 marcas", saida)
        doc = parte(sai)
        self.assertEqual(contar(doc, "r"), 1)
        self.assertEqual(textos_dos_paragrafos(doc), ["Olá mundo cruel"])
        self.assert_mesmo_texto_e_formato(ent, sai)

    def test_ordem_das_propriedades_nao_importa(self):
        corpo = "<w:p>" + r("a", "<w:b/><w:i/>") + r("b", "<w:i/><w:b/>") + "</w:p>"
        ent, sai, _ = self.juntar(corpo)
        self.assertEqual(contar(parte(sai), "r"), 1)

    def test_nao_junta_formatacao_diferente(self):
        corpo = ("<w:p>" + r("negrito", "<w:b/>") + r("itálico", "<w:i/>") + r("normal")
                 + r("cor", '<w:color w:val="FF0000"/>') + r("cor2", '<w:color w:val="00FF00"/>') + "</w:p>")
        ent, sai, saida = self.juntar(corpo)
        self.assertIn("Juntei 0 trechos", saida)
        self.assertEqual(contar(parte(sai), "r"), 5)
        self.assert_mesmo_texto_e_formato(ent, sai)

    def test_nao_atravessa_marcas_campos_notas_quebras(self):
        corpo = ("<w:p>"
                 + r("A") + '<w:commentRangeStart w:id="5"/>' + r("B")
                 + '<w:bookmarkStart w:id="1" w:name="m"/>' + r("C") + '<w:bookmarkEnd w:id="1"/>'
                 + '<w:ins w:id="2" w:author="X" w:date="2026-01-01T00:00:00Z">' + r("D") + "</w:ins>"
                 + r("E")
                 + '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                 + '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
                 + '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                 + r("1")
                 + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
                 + r("F") + '<w:r><w:footnoteReference w:id="1"/></w:r>' + r("G")
                 + '<w:r><w:br/></w:r>' + r("H")
                 + '<w:r><w:t>I</w:t><w:br w:type="page"/></w:r>' + r("J")
                 + '<w:commentRangeEnd w:id="5"/><w:r><w:commentReference w:id="5"/></w:r>'
                 + "</w:p>")
        ent, sai, saida = self.juntar(corpo)
        self.assertIn("Juntei 0 trechos", saida)
        self.assertEqual(contar(parte(ent), "r"), contar(parte(sai), "r"))
        self.assert_mesmo_texto_e_formato(ent, sai)

    def test_junta_dentro_do_mesmo_conteiner(self):
        corpo = ("<w:p>" + r("fora") + '<w:ins w:id="1" w:author="X" w:date="2026-01-01T00:00:00Z">'
                 + r("den") + r("tro") + "</w:ins>"
                 + '<w:del w:id="2" w:author="X" w:date="2026-01-01T00:00:00Z">'
                 + '<w:r><w:delText xml:space="preserve">ex </w:delText></w:r>'
                 + '<w:r><w:delText>cluído</w:delText></w:r></w:del>' + r("fim") + "</w:p>")
        ent, sai, saida = self.juntar(corpo)
        self.assertIn("Juntei 2 trechos", saida)
        doc = parte(sai)
        ins = doc.find(".//" + w("ins"))
        self.assertEqual(len(ins.findall(w("r"))), 1)
        dele = doc.find(".//" + w("del"))
        self.assertEqual(dele.find(".//" + w("delText")).text, "ex cluído")
        self.assert_mesmo_texto_e_formato(ent, sai)

    def test_preserva_espaco_ao_juntar(self):
        corpo = "<w:p>" + r("um ") + r("dois", preservar=False) + r("  três  ") + "</w:p>"
        ent, sai, _ = self.juntar(corpo)
        t = parte(sai).find(".//" + w("t"))
        self.assertEqual(t.text, "um dois  três  ")
        self.assertEqual(t.get(XML_SPACE), "preserve")
        self.assert_mesmo_texto_e_formato(ent, sai)

    def test_espaco_sem_preserve_na_ponta_nao_junta(self):
        corpo = "<w:p>" + r("um ", preservar=False) + r("dois") + "</w:p>"
        ent, sai, saida = self.juntar(corpo)
        self.assertIn("Juntei 0 trechos", saida)

    def test_tabulacao_e_hifen_juntam(self):
        corpo = ('<w:p><w:r><w:t>a</w:t><w:tab/></w:r><w:r><w:t>b</w:t><w:noBreakHyphen/></w:r>'
                 + r("c") + "</w:p>")
        ent, sai, saida = self.juntar(corpo)
        self.assertIn("Juntei 2 trechos", saida)
        run = parte(sai).find(".//" + w("r"))
        self.assertEqual([c.tag.split("}")[1] for c in run], ["t", "tab", "t", "noBreakHyphen", "t"])

    def test_cabecalho_rodape_e_notas(self):
        cab = cabecalho_xml("<w:p>" + r("Cabe", rsid="1") + r("çalho", rsid="2") + "</w:p>")
        rod = cabecalho_xml("<w:p>" + r("Roda", rsid="1") + r("pé", rsid="2") + "</w:p>", "ftr")
        notas = cabecalho_xml(
            '<w:footnote w:id="1"><w:p>' + r("No") + r("ta") + "</w:p></w:footnote>", "footnotes")
        extras = [
            ("word/header1.xml", cab, CT_BASE + "header+xml", REL_BASE + "header", "rIdH1"),
            ("word/footer1.xml", rod, CT_BASE + "footer+xml", REL_BASE + "footer", "rIdF1"),
            ("word/footnotes.xml", notas, CT_BASE + "footnotes+xml", REL_BASE + "footnotes", "rIdN1"),
        ]
        sect = ('<w:headerReference w:type="default" r:id="rIdH1"/>'
                '<w:footerReference w:type="default" r:id="rIdF1"/>')
        ent, sai, saida = self.juntar("<w:p>" + r("corpo") + "</w:p>", extras, sect)
        self.assertIn("Juntei 3 trechos", saida)
        for nome in ("word/header1.xml", "word/footer1.xml", "word/footnotes.xml"):
            self.assertEqual(contar(parte(sai, nome), "r"), 1, nome)
            self.assert_mesmo_texto_e_formato(ent, sai, nome)

    def test_partes_nao_mexidas_ficam_identicas(self):
        ent, sai, _ = self.juntar("<w:p>" + r("a") + r("b") + "</w:p>")
        for nome in ("word/styles.xml", "word/settings.xml", "word/numbering.xml"):
            self.assertEqual(bytes_da_parte(ent, nome), bytes_da_parte(sai, nome), nome)
        self.assertEqual(nomes(sai)[0], "[Content_Types].xml")
        self.assertEqual(sorted(nomes(ent)), sorted(nomes(sai)))

    def test_documento_feito_no_python_docx(self):
        d = docx.Document()
        p = d.add_paragraph()
        for pedaco in ["Este ", "texto ", "veio ", "em ", "pedaços."]:
            p.add_run(pedaco).bold = True
        p.add_run(" Fim.").italic = True
        ent = self.arq("pd.docx")
        d.save(ent)
        sai = self.arq("pd_sai.docx")
        proc = rodar("juntar_trechos.py", ent, "-o", sai)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertIn("Juntei 4 trechos", proc.stdout)
        lido = self.abre_no_python_docx(sai)
        self.assertEqual(lido.paragraphs[0].text, "Este texto veio em pedaços. Fim.")
        self.assertEqual(len(lido.paragraphs[0].runs), 2)
        self.assert_mesmo_texto_e_formato(ent, sai)

    def test_pasta_desempacotada_no_lugar(self):
        ent = montar(self.arq("ent.docx"), "<w:p>" + r("a", rsid="1") + r("b", rsid="2") + "</w:p>")
        pasta = self.arq("desemp")
        with zipfile.ZipFile(ent) as z:
            z.extractall(pasta)
        proc = rodar("juntar_trechos.py", pasta)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        with open(os.path.join(pasta, "word", "document.xml"), "rb") as f:
            self.assertEqual(contar(etree.fromstring(f.read()), "r"), 1)

    def test_erros_de_entrada(self):
        self.assertEqual(rodar("juntar_trechos.py", self.arq("nao_existe.docx"), "-o", self.arq("x.docx")).returncode, 3)
        falso = self.arq("falso.docx")
        with open(falso, "w") as f:
            f.write("não sou zip")
        self.assertEqual(rodar("juntar_trechos.py", falso, "-o", self.arq("x.docx")).returncode, 3)
        ent = montar(self.arq("ent.docx"), "<w:p>" + r("a") + "</w:p>")
        self.assertEqual(rodar("juntar_trechos.py", ent).returncode, 2)
        self.assertEqual(rodar("juntar_trechos.py", "-h").returncode, 0)


# ====================================================================== comentar

def texto_comentado(doc, id_):
    """Texto entre commentRangeStart e commentRangeEnd do id dado, em ordem de documento."""
    dentro = False
    pedacos = []
    for el in doc.iter():
        if el.tag == w("commentRangeStart") and el.get(w("id")) == str(id_):
            dentro = True
        elif el.tag == w("commentRangeEnd") and el.get(w("id")) == str(id_):
            dentro = False
        elif dentro and el.tag == w("t"):
            pedacos.append(el.text or "")
    return "".join(pedacos)


def id_criado(saida):
    return int(re.search(r"Comentário (\d+) criado", saida).group(1))


class TestComentar(Base):
    CORPO = ("<w:p>" + r("O prazo ") + r("de entrega", "<w:b/>") + r(" é curto demais.") + "</w:p>"
             "<w:p>" + r("Segundo parágrafo sobre o prazo de pagamento.") + "</w:p>")

    def comentar(self, ent, *args):
        sai = self.arq("com_%d.docx" % len(os.listdir(self.pasta)))
        proc = rodar("comentar.py", ent, *args, "-o", sai)
        return proc, sai

    def test_trecho_que_atravessa_formatacoes(self):
        ent = montar(self.arq("ent.docx"), self.CORPO)
        proc, sai = self.comentar(ent, "--trecho", "prazo de entrega é", "--texto", "Curto mesmo",
                                  "--autor", "Ana Souza")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        cid = id_criado(proc.stdout)
        doc = parte(sai)
        self.assertEqual(texto_comentado(doc, cid), "prazo de entrega é")
        self.assertEqual(textos_dos_paragrafos(doc), textos_dos_paragrafos(parte(ent)))
        # formatação por caractere igual
        self.assertEqual([(c, f) for c, f, _ in formato_por_caractere(doc)],
                         [(c, f) for c, f, _ in formato_por_caractere(parte(ent))])
        # partes, relações e tipos
        for nome in ("word/comments.xml", "word/commentsExtended.xml", "word/commentsIds.xml"):
            self.assertIn(nome, nomes(sai))
        ct = etree.tostring(parte(sai, "[Content_Types].xml")).decode()
        self.assertIn("/word/comments.xml", ct)
        self.assertIn("comments+xml", ct)
        self.assertIn("commentsExtended+xml", ct)
        self.assertIn("commentsIds+xml", ct)
        rels = etree.tostring(parte(sai, "word/_rels/document.xml.rels")).decode()
        self.assertIn("relationships/comments", rels)
        self.assertIn("commentsExtended", rels)
        self.assertIn("commentsIds", rels)
        com = parte(sai, "word/comments.xml").find(w("comment"))
        self.assertEqual(com.get(w("author")), "Ana Souza")
        self.assertEqual(com.get(w("initials")), "AS")
        self.assertEqual(contar(doc, "commentReference"), 1)
        # ordem: início, trechos, fim, referência
        p = doc.find(w("body")).find(w("p"))
        tags = [c.tag.split("}")[1] for c in p]
        self.assertLess(tags.index("commentRangeStart"), tags.index("commentRangeEnd"))
        fim = p.find(w("commentRangeEnd"))
        self.assertIsNotNone(fim.getnext().find(w("commentReference")))
        # python-docx lê o comentário
        lido = self.abre_no_python_docx(sai)
        comentarios = list(lido.comments)
        self.assertEqual(len(comentarios), 1)
        self.assertEqual(comentarios[0].text, "Curto mesmo")
        self.assertEqual(comentarios[0].author, "Ana Souza")

    def test_trecho_no_meio_de_um_trecho_divide(self):
        ent = montar(self.arq("ent.docx"),
                     "<w:p>" + r("Começo meio final", '<w:i/><w:color w:val="112233"/>') + "</w:p>")
        proc, sai = self.comentar(ent, "--trecho", "meio", "--texto", "Olha aqui")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        doc = parte(sai)
        runs_texto = [x for x in doc.iter(w("r")) if x.find(w("t")) is not None]
        self.assertEqual([x.find(w("t")).text for x in runs_texto], ["Começo ", "meio", " final"])
        for x in runs_texto:
            self.assertEqual(assinatura_rpr(x), assinatura_rpr(parte(ent).find(".//" + w("r"))))
        self.assertEqual(runs_texto[0].find(w("t")).get(XML_SPACE), "preserve")
        self.assertEqual(texto_comentado(doc, id_criado(proc.stdout)), "meio")

    def test_trecho_inexistente_sai_1(self):
        ent = montar(self.arq("ent.docx"), self.CORPO)
        proc, sai = self.comentar(ent, "--trecho", "não está lá", "--texto", "x")
        self.assertEqual(proc.returncode, 1)
        self.assertIn("não encontrado", proc.stderr)
        self.assertFalse(os.path.exists(sai))

    def test_trecho_ambiguo_e_ocorrencia(self):
        ent = montar(self.arq("ent.docx"), self.CORPO)
        proc, _ = self.comentar(ent, "--trecho", "prazo", "--texto", "x")
        self.assertEqual(proc.returncode, 1)
        self.assertIn("aparece 2 vezes", proc.stderr)
        self.assertIn("--ocorrencia", proc.stderr)
        proc, sai = self.comentar(ent, "--trecho", "prazo", "--ocorrencia", "2", "--texto", "segundo")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        doc = parte(sai)
        ps = doc.find(w("body")).findall(w("p"))
        self.assertIsNone(ps[0].find(w("commentRangeStart")))
        self.assertIsNotNone(ps[1].find(w("commentRangeStart")))
        self.assertEqual(texto_comentado(doc, id_criado(proc.stdout)), "prazo")
        proc, _ = self.comentar(ent, "--trecho", "prazo", "--ocorrencia", "3", "--texto", "x")
        self.assertEqual(proc.returncode, 1)
        proc, _ = self.comentar(ent, "--trecho", "prazo", "--ocorrencia", "0", "--texto", "x")
        self.assertEqual(proc.returncode, 2)

    def test_resposta_encadeada(self):
        ent = montar(self.arq("ent.docx"), self.CORPO)
        proc, um = self.comentar(ent, "--trecho", "curto demais", "--texto", "Pai", "--autor", "Ana")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        pai = id_criado(proc.stdout)
        proc, dois = self.comentar(um, "--resposta-a", str(pai), "--texto", "Filho\nsegunda linha",
                                   "--autor", "Bia")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        filho = id_criado(proc.stdout)
        self.assertNotEqual(pai, filho)
        doc = parte(dois)
        self.assertEqual(texto_comentado(doc, filho), "curto demais")
        self.assertEqual(contar(doc, "commentReference"), 2)
        coms = parte(dois, "word/comments.xml")
        por_id = {c.get(w("id")): c for c in coms.iter(w("comment"))}
        self.assertEqual(len(por_id[str(filho)].findall(w("p"))), 2)
        para_pai = por_id[str(pai)].findall(w("p"))[-1].get("{%s}paraId" % NS["w14"])
        para_filho = por_id[str(filho)].findall(w("p"))[-1].get("{%s}paraId" % NS["w14"])
        ext = parte(dois, "word/commentsExtended.xml")
        entradas = {e.get("{%s}paraId" % NS["w15"]): e for e in ext}
        self.assertIn(para_pai, entradas)
        self.assertEqual(entradas[para_filho].get("{%s}paraIdParent" % NS["w15"]), para_pai)
        self.assertIsNone(entradas[para_pai].get("{%s}paraIdParent" % NS["w15"]))
        ids = parte(dois, "word/commentsIds.xml")
        duraveis = [e.get("{%s}durableId" % NS["w16cid"]) for e in ids]
        self.assertEqual(len(duraveis), 2)
        self.assertEqual(len(set(duraveis)), 2)
        # rel e tipo não duplicados
        rels = parte(dois, "word/_rels/document.xml.rels")
        alvos = [x.get("Target") for x in rels]
        self.assertEqual(alvos.count("comments.xml"), 1)
        self.assertEqual(alvos.count("commentsExtended.xml"), 1)
        lido = self.abre_no_python_docx(dois)
        self.assertEqual(len(list(lido.comments)), 2)

    def test_resposta_a_comentario_inexistente(self):
        ent = montar(self.arq("ent.docx"), self.CORPO)
        proc, _ = self.comentar(ent, "--resposta-a", "7", "--texto", "x")
        self.assertEqual(proc.returncode, 1)
        self.assertIn("Não existe comentário", proc.stderr)

    def test_soma_a_comentario_ja_existente(self):
        d = docx.Document()
        p = d.add_paragraph()
        run = p.add_run("Texto com comentário antigo e outro trecho.")
        d.add_comment(run, text="Antigo", author="Velho", initials="V")
        ent = self.arq("existente.docx")
        d.save(ent)
        proc, sai = self.comentar(ent, "--trecho", "outro trecho", "--texto", "Novo")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        coms = parte(sai, "word/comments.xml")
        ids = [c.get(w("id")) for c in coms.iter(w("comment"))]
        self.assertEqual(len(ids), 2)
        self.assertEqual(len(set(ids)), 2)
        self.assertNotIn("ns0", etree.tostring(coms).decode())
        lido = self.abre_no_python_docx(sai)
        self.assertEqual(sorted(c.text for c in lido.comments), ["Antigo", "Novo"])
        proc, sai2 = self.comentar(sai, "--resposta-a", ids[0], "--texto", "Resposta ao antigo")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertEqual(len(list(self.abre_no_python_docx(sai2).comments)), 3)

    def test_autor_padrao_pela_variavel(self):
        ent = montar(self.arq("ent.docx"), self.CORPO)
        sai = self.arq("var.docx")
        env = dict(os.environ, DOC_AUTOR="Carla Dias Lima")
        proc = subprocess.run([sys.executable, os.path.join(SCRIPTS, "comentar.py"), ent,
                               "--trecho", "Segundo", "--texto", "x", "-o", sai],
                              capture_output=True, text=True, env=env)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        com = parte(sai, "word/comments.xml").find(w("comment"))
        self.assertEqual(com.get(w("author")), "Carla Dias Lima")
        self.assertEqual(com.get(w("initials")), "CDL")

    def test_argumentos(self):
        ent = montar(self.arq("ent.docx"), self.CORPO)
        self.assertEqual(rodar("comentar.py", ent, "--texto", "x", "-o", self.arq("a.docx")).returncode, 2)
        self.assertEqual(rodar("comentar.py", ent, "--trecho", "prazo de", "--texto", "x").returncode, 2)
        self.assertEqual(rodar("comentar.py", "-h").returncode, 0)

    def test_abre_no_libreoffice(self):
        ent = montar(self.arq("ent.docx"), self.CORPO)
        proc, um = self.comentar(ent, "--trecho", "de entrega", "--texto", "Pai")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        proc, dois = self.comentar(um, "--resposta-a", str(id_criado(proc.stdout)), "--texto", "Filho")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        texto = self.abre_no_libreoffice(dois)
        if texto:
            self.assertIn("prazo", texto)
        # O LibreOffice reconhece o encadeamento: a resposta aponta para o pai.
        perfil = tempfile.mkdtemp(prefix="perfil_lo_", dir=self.pasta)
        odt_dir = os.path.join(self.pasta, "odt")
        subprocess.run([SOFFICE, "-env:UserInstallation=file://" + perfil, "--headless",
                        "--convert-to", "odt", "--outdir", odt_dir, dois],
                       capture_output=True, timeout=180)
        odt = os.path.join(odt_dir, os.path.splitext(os.path.basename(dois))[0] + ".odt")
        with zipfile.ZipFile(odt) as z:
            conteudo = z.read("content.xml").decode("utf-8")
        self.assertEqual(conteudo.count("<office:annotation "), 2)
        self.assertIn("parent-name=", conteudo)


# ====================================================================== aceitar

def ins(conteudo, id_="1"):
    return '<w:ins w:id="%s" w:author="Revisor" w:date="2026-01-01T00:00:00Z">%s</w:ins>' % (id_, conteudo)


def dele(texto, id_="2"):
    return ('<w:del w:id="%s" w:author="Revisor" w:date="2026-01-01T00:00:00Z">'
            '<w:r><w:delText xml:space="preserve">%s</w:delText></w:r></w:del>' % (id_, texto))


MARCA_DEL = '<w:del w:id="90" w:author="Revisor" w:date="2026-01-01T00:00:00Z"/>'
MARCA_INS = '<w:ins w:id="91" w:author="Revisor" w:date="2026-01-01T00:00:00Z"/>'
CHG = 'w:author="Revisor" w:date="2026-01-01T00:00:00Z"'
PROIBIDOS = ["ins", "del", "moveFrom", "moveTo", "delText", "delInstrText", "rPrChange",
             "pPrChange", "sectPrChange", "tblPrChange", "tblPrExChange", "trPrChange",
             "tcPrChange", "tblGridChange", "numberingChange", "moveFromRangeStart",
             "moveFromRangeEnd", "moveToRangeStart", "moveToRangeEnd", "cellIns", "cellDel"]


def celula(texto, tcpr=""):
    return "<w:tc><w:tcPr><w:tcW w:w=\"2000\" w:type=\"dxa\"/>%s</w:tcPr><w:p>%s</w:p></w:tc>" % (
        tcpr, r(texto))


class TestAceitarAlteracoes(Base):
    def aceitar(self, corpo, extras=(), sect_extra="", *opcoes):
        ent = montar(self.arq("ent.docx"), corpo, extras, sect_extra)
        sai = self.arq("sai.docx")
        proc = rodar("aceitar_alteracoes.py", ent, "-o", sai, *opcoes)
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        return ent, sai, proc.stdout

    def assert_sem_marcas(self, caminho, nome="word/document.xml"):
        raiz = parte(caminho, nome)
        for marca in PROIBIDOS:
            self.assertEqual(contar(raiz, marca), 0, "sobrou w:%s em %s" % (marca, nome))

    def test_insercao_e_exclusao(self):
        corpo = ("<w:p>" + r("Era ") + ins(r("muito ")) + dele("pouco ") + r("bom.")
                 + ins(r("!", "<w:b/>"), "3") + "</w:p>")
        ent, sai, saida = self.aceitar(corpo)
        doc = parte(sai)
        self.assertEqual(textos_dos_paragrafos(doc), ["Era muito bom.!"])
        self.assertIn("inserções: 2", saida)
        self.assertIn("exclusões: 1", saida)
        self.assert_sem_marcas(sai)
        self.assertEqual(self.abre_no_python_docx(sai).paragraphs[0].text, "Era muito bom.!")

    def test_movimento(self):
        corpo = ('<w:p><w:moveFromRangeStart w:id="10" w:name="mv1" %s/>' % CHG
                 + '<w:moveFrom w:id="11" %s><w:r><w:t>Frase movida.</w:t></w:r></w:moveFrom>' % CHG
                 + '<w:moveFromRangeEnd w:id="10"/>' + r(" Fica.") + "</w:p>"
                 + "<w:p>" + r("Outro. ")
                 + '<w:moveToRangeStart w:id="12" w:name="mv1" %s/>' % CHG
                 + '<w:moveTo w:id="13" %s><w:r><w:t>Frase movida.</w:t></w:r></w:moveTo>' % CHG
                 + '<w:moveToRangeEnd w:id="12"/></w:p>')
        ent, sai, _ = self.aceitar(corpo)
        self.assertEqual(textos_dos_paragrafos(parte(sai)), [" Fica.", "Outro. Frase movida."])
        self.assert_sem_marcas(sai)

    def test_mudancas_de_formatacao_ficam_no_estado_novo(self):
        corpo = ('<w:p><w:pPr><w:jc w:val="center"/><w:pPrChange w:id="20" %s><w:pPr><w:jc w:val="left"/>'
                 '</w:pPr></w:pPrChange></w:pPr>'
                 '<w:r><w:rPr><w:b/><w:rPrChange w:id="21" %s><w:rPr><w:i/></w:rPr></w:rPrChange></w:rPr>'
                 '<w:t>Negrito novo</w:t></w:r></w:p>' % (CHG, CHG)
                 + '<w:tbl><w:tblPr><w:tblW w:w="4000" w:type="dxa"/><w:tblPrChange w:id="22" %s>'
                   '<w:tblPr><w:tblW w:w="3000" w:type="dxa"/></w:tblPr></w:tblPrChange></w:tblPr>' % CHG
                 + '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/><w:tblGridChange w:id="23">'
                   '<w:tblGrid><w:gridCol w:w="1500"/><w:gridCol w:w="1500"/></w:tblGrid></w:tblGridChange></w:tblGrid>'
                 + '<w:tr><w:trPr><w:jc w:val="center"/><w:trPrChange w:id="24" %s><w:trPr/></w:trPrChange></w:trPr>' % CHG
                 + celula("A", '<w:tcPrChange w:id="25" %s><w:tcPr/></w:tcPrChange>' % CHG)
                 + celula("B") + "</w:tr></w:tbl><w:p/>")
        sect = '<w:sectPrChange w:id="26" %s><w:sectPr/></w:sectPrChange>' % CHG
        ent, sai, saida = self.aceitar(corpo, (), sect)
        doc = parte(sai)
        self.assertIn("mudanças de formatação: 7", saida)
        self.assert_sem_marcas(sai)
        self.assertIsNotNone(doc.find(".//" + w("r") + "/" + w("rPr") + "/" + w("b")))
        self.assertIsNone(doc.find(".//" + w("i")))
        self.assertEqual(doc.find(".//" + w("jc")).get(w("val")), "center")
        self.assertEqual(doc.find(".//" + w("tblW")).get(w("w")), "4000")
        self.assertEqual([g.get(w("w")) for g in doc.iter(w("gridCol"))], ["2000", "2000"])
        self.assertIsNotNone(doc.find(".//" + w("sectPr") + "/" + w("pgSz")))

    def test_linhas_e_celulas_de_tabela(self):
        grade = '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
        corpo = ('<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>' + grade
                 + "<w:tr>" + celula("fica1") + celula("fica2") + "</w:tr>"
                 + "<w:tr><w:trPr>" + MARCA_INS + "</w:trPr>" + celula("nova1") + celula("nova2") + "</w:tr>"
                 + "<w:tr><w:trPr>" + MARCA_DEL + "</w:trPr>" + celula("sai1") + celula("sai2") + "</w:tr>"
                 + "<w:tr>" + celula("c1") + celula("c2", '<w:cellDel w:id="30" %s/>' % CHG) + "</w:tr>"
                 + "</w:tbl><w:p/>"
                 + '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>' + grade
                 + "<w:tr><w:trPr>" + MARCA_DEL + "</w:trPr>" + celula("x") + celula("y") + "</w:tr></w:tbl>"
                 + "<w:p>" + r("depois") + "</w:p>")
        ent, sai, saida = self.aceitar(corpo)
        doc = parte(sai)
        textos = [t for t in textos_dos_paragrafos(doc) if t]
        self.assertEqual(textos, ["fica1", "fica2", "nova1", "nova2", "c1", "depois"])
        self.assertEqual(contar(doc, "tbl"), 1)
        self.assertIn("linhas excluídas: 2", saida)
        self.assertIn("tabelas que ficaram vazias: 1", saida)
        self.assert_sem_marcas(sai)
        self.abre_no_python_docx(sai)

    def test_marca_de_paragrafo_excluida(self):
        corpo = ("<w:p><w:pPr><w:rPr>" + MARCA_DEL + "</w:rPr></w:pPr>" + r("Primeira ") + "</w:p>"
                 + '<w:p><w:pPr><w:jc w:val="right"/></w:pPr>' + r("segunda.") + "</w:p>"
                 # parágrafo apagado inteiro
                 + "<w:p><w:pPr><w:rPr>" + MARCA_DEL + "</w:rPr></w:pPr>" + dele("Apagado inteiro.") + "</w:p>"
                 + "<w:p>" + r("Terceira.") + "</w:p>"
                 # parágrafo inserido inteiro
                 + "<w:p><w:pPr><w:rPr>" + MARCA_INS + "</w:rPr></w:pPr>" + ins(r("Novo.")) + "</w:p>"
                 # duas marcas seguidas
                 + "<w:p><w:pPr><w:rPr>" + MARCA_DEL + "</w:rPr></w:pPr>" + r("a") + "</w:p>"
                 + "<w:p><w:pPr><w:rPr>" + MARCA_DEL + "</w:rPr></w:pPr>" + r("b") + "</w:p>"
                 + "<w:p>" + r("c") + "</w:p>")
        ent, sai, saida = self.aceitar(corpo)
        doc = parte(sai)
        self.assertEqual(textos_dos_paragrafos(doc), ["Primeira segunda.", "Terceira.", "Novo.", "abc"])
        primeiro = doc.find(w("body")).find(w("p"))
        self.assertEqual(primeiro.find(".//" + w("jc")).get(w("val")), "right")
        self.assertIn("parágrafos juntados: 4", saida)
        self.assert_sem_marcas(sai)

    def test_paragrafo_excluido_antes_de_tabela(self):
        grade = '<w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid>'
        corpo = ("<w:p><w:pPr><w:rPr>" + MARCA_DEL + "</w:rPr></w:pPr>" + dele("some") + "</w:p>"
                 + "<w:tbl><w:tblPr/>" + grade + "<w:tr>" + celula("t") + "</w:tr></w:tbl><w:p/>")
        ent, sai, _ = self.aceitar(corpo)
        corpo_sai = parte(sai).find(w("body"))
        self.assertEqual(corpo_sai[0].tag, w("tbl"))

    def test_cabecalho_rodape_e_notas(self):
        cab = cabecalho_xml("<w:p>" + r("Cab ") + ins(r("novo")) + dele(" velho") + "</w:p>")
        rod = cabecalho_xml("<w:p>" + dele("Tirar ") + r("Rodapé") + "</w:p>", "ftr")
        notas = cabecalho_xml(
            '<w:footnote w:id="1"><w:p>' + r("Nota ") + ins(r("certa")) + dele("errada") + "</w:p></w:footnote>",
            "footnotes")
        fim = cabecalho_xml(
            '<w:endnote w:id="1"><w:p>' + dele("x") + r("Final") + "</w:p></w:endnote>", "endnotes")
        extras = [
            ("word/header1.xml", cab, CT_BASE + "header+xml", REL_BASE + "header", "rIdH1"),
            ("word/footer1.xml", rod, CT_BASE + "footer+xml", REL_BASE + "footer", "rIdF1"),
            ("word/footnotes.xml", notas, CT_BASE + "footnotes+xml", REL_BASE + "footnotes", "rIdN1"),
            ("word/endnotes.xml", fim, CT_BASE + "endnotes+xml", REL_BASE + "endnotes", "rIdE1"),
        ]
        sect = ('<w:headerReference w:type="default" r:id="rIdH1"/>'
                '<w:footerReference w:type="default" r:id="rIdF1"/>')
        corpo = "<w:p>" + r("Texto") + '<w:r><w:footnoteReference w:id="1"/></w:r></w:p>'
        ent, sai, _ = self.aceitar(corpo, extras, sect)
        esperado = {
            "word/header1.xml": ["Cab novo"],
            "word/footer1.xml": ["Rodapé"],
            "word/footnotes.xml": ["Nota certa"],
            "word/endnotes.xml": ["Final"],
        }
        for nome, textos in esperado.items():
            self.assertEqual(textos_dos_paragrafos(parte(sai, nome)), textos, nome)
            self.assert_sem_marcas(sai, nome)

    def test_texto_excluido_solto_e_campo_excluido(self):
        corpo = ("<w:p>" + r("A")
                 + '<w:del w:id="40" %s><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                   '<w:r><w:delInstrText> PAGE </w:delInstrText></w:r>'
                   '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:del>' % CHG
                 + '<w:r><w:delText>solto</w:delText></w:r>' + r("B") + "</w:p>")
        ent, sai, saida = self.aceitar(corpo)
        self.assertEqual(textos_dos_paragrafos(parte(sai)), ["AB"])
        self.assertIn("texto excluído solto: 1", saida)
        self.assertEqual(contar(parte(sai), "fldChar"), 0)
        self.assert_sem_marcas(sai)

    def test_desligar_rastreio(self):
        ent = montar(self.arq("ent.docx"), "<w:p>" + ins(r("x")) + "</w:p>")
        pasta = self.arq("desemp")
        with zipfile.ZipFile(ent) as z:
            z.extractall(pasta)
        cfg = os.path.join(pasta, "word", "settings.xml")
        with open(cfg, "rb") as f:
            raiz = etree.fromstring(f.read())
        etree.SubElement(raiz, w("trackRevisions"))
        with open(cfg, "wb") as f:
            f.write(etree.tostring(raiz, xml_declaration=True, encoding="UTF-8"))
        sai = self.arq("sai.docx")
        proc = rodar("aceitar_alteracoes.py", pasta, "-o", sai, "--desligar-rastreio")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertEqual(contar(parte(sai, "word/settings.xml"), "trackRevisions"), 0)
        self.assertEqual(textos_dos_paragrafos(parte(sai)), ["x"])

    def test_sem_alteracoes_nao_mexe(self):
        ent, sai, saida = self.aceitar("<w:p>" + r("nada") + "</w:p>")
        self.assertIn("Não havia alteração rastreada", saida)
        self.assertEqual(bytes_da_parte(ent, "word/document.xml"), bytes_da_parte(sai, "word/document.xml"))

    def test_abre_no_python_docx_e_libreoffice(self):
        corpo = ("<w:p>" + r("Contrato ") + ins(r("revisado ")) + dele("antigo ") + r("final.") + "</w:p>"
                 + "<w:p><w:pPr><w:rPr>" + MARCA_DEL + "</w:rPr></w:pPr>" + r("Junta ") + "</w:p>"
                 + "<w:p>" + r("aqui.") + "</w:p>")
        ent, sai, _ = self.aceitar(corpo)
        lido = self.abre_no_python_docx(sai)
        self.assertEqual([p.text for p in lido.paragraphs], ["Contrato revisado final.", "Junta aqui."])
        texto = self.abre_no_libreoffice(sai)
        if texto:
            self.assertIn("Contrato revisado final.", texto)
            self.assertNotIn("antigo", texto)

    def test_funcao_de_sobras(self):
        import aceitar_alteracoes
        raiz = etree.fromstring('<w:body %s><w:p><w:r><w:delText>x</w:delText></w:r></w:p></w:body>' % DECL)
        self.assertEqual(aceitar_alteracoes.sobras(raiz), {"delText": 1})


if __name__ == "__main__":
    unittest.main()
