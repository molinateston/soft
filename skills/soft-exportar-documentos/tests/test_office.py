"""Testes dos scripts de office/ (desempacotar, empacotar, conferir, soffice_seguro).

Rodar de dentro de novo/:  python3 -m unittest tests.test_office -v
Os arquivos de teste são gerados aqui mesmo com python-docx, python-pptx e openpyxl.
"""

import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(AQUI), "scripts", "office")
sys.path.insert(0, SCRIPTS)

import docx  # noqa: E402
import openpyxl  # noqa: E402
import pptx  # noqa: E402
from lxml import etree  # noqa: E402

import comum  # noqa: E402
import conferir  # noqa: E402
import soffice_seguro  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
DATA = "2026-10-15T10:00:00Z"


def rodar(script, *args):
    proc = subprocess.run([sys.executable, os.path.join(SCRIPTS, script)] + [str(a) for a in args],
                          capture_output=True, text=True, timeout=300)
    return proc.returncode, proc.stdout + proc.stderr


def gerar_docx(caminho, paragrafos=("Olá mundo", "Segundo parágrafo")):
    d = docx.Document()
    for p in paragrafos:
        d.add_paragraph(p)
    d.save(caminho)


def gerar_pptx(caminho, slides=2):
    pr = pptx.Presentation()
    for i in range(slides):
        s = pr.slides.add_slide(pr.slide_layouts[1])
        s.shapes.title.text = "Título %d" % (i + 1)
    pr.save(caminho)


def gerar_xlsx(caminho):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws["A1"] = "nome"
    ws["B1"] = "=1+1"
    wb.create_sheet("Outra")["A1"] = "texto"
    wb.save(caminho)


def ler_parte(caminho, parte):
    with zipfile.ZipFile(caminho) as zf:
        return zf.read(parte)


def trocar_partes(origem, destino, trocas=None, apagar=(), extras=None):
    """Copia o zip trocando, apagando ou acrescentando partes. trocas: {parte: bytes ou função}."""
    trocas = trocas or {}
    with zipfile.ZipFile(origem) as zi, zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zi.infolist():
            if info.filename in apagar:
                continue
            dados = zi.read(info)
            if info.filename in trocas:
                t = trocas[info.filename]
                dados = t(dados) if callable(t) else t
            zo.writestr(info.filename, dados)
        for nome, dados in (extras or {}).items():
            zo.writestr(nome, dados)


def editar_documento(dados, funcao):
    raiz = etree.fromstring(dados)
    funcao(raiz)
    return etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)


def fragmento(xml):
    raiz = etree.fromstring('<w:x xmlns:w="%s">%s</w:x>' % (W, xml))
    return list(raiz)


def paragrafos(raiz):
    return raiz.findall(".//{%s}body/{%s}p" % (W, W))


def substituir_conteudo(par, xml):
    for filho in list(par):
        if etree.QName(filho).localname != "pPr":
            par.remove(filho)
    for novo in fragmento(xml):
        par.append(novo)


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="teste-office-")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def p(self, *nomes):
        return os.path.join(self.tmp, *nomes)


class TestAjuda(Base):
    def test_todos_respondem_ajuda(self):
        for s in ("desempacotar.py", "empacotar.py", "conferir.py", "soffice_seguro.py"):
            codigo, saida = rodar(s, "-h")
            self.assertEqual(codigo, 0, s)
            self.assertIn("usage", saida.lower())

    def test_argumento_errado_sai_2(self):
        self.assertEqual(rodar("conferir.py")[0], 2)
        self.assertEqual(rodar("empacotar.py", "so-um")[0], 2)


class TestRecuo(unittest.TestCase):
    def ida_e_volta(self, xml):
        arv = comum.ler_xml(xml)
        comum.recuar(arv.getroot())
        recuado = comum.gravar_xml(arv, xml)
        arv2 = comum.ler_xml(recuado)
        comum.tirar_recuo(arv2.getroot())
        return recuado, comum.gravar_xml(arv2, recuado)

    def test_texto_de_folha_e_preserve_intactos(self):
        xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
               '<w:document xmlns:w="%s"><w:body><w:p><w:r><w:t xml:space="preserve"> </w:t></w:r>'
               '<w:r><w:t>  dois  </w:t></w:r></w:p></w:body></w:document>' % W).encode()
        recuado, volta = self.ida_e_volta(xml)
        self.assertIn(b"\n    <w:p>", recuado)
        self.assertIn(b'<w:t xml:space="preserve"> </w:t>', recuado)
        self.assertIn(b"<w:t>  dois  </w:t>", recuado)
        self.assertEqual(volta, xml)

    def test_conteudo_misto_nao_mexe(self):
        xml = b'<raiz><p>texto <b>negrito</b> fim</p><q><r/></q></raiz>'
        recuado, volta = self.ida_e_volta(xml)
        self.assertIn(b"<p>texto <b>negrito</b> fim</p>", recuado)
        self.assertEqual(volta, xml)

    def test_preserve_herdado(self):
        xml = b'<raiz xml:space="preserve"><a> <b/> </a></raiz>'
        recuado, volta = self.ida_e_volta(xml)
        self.assertEqual(recuado, xml)
        self.assertEqual(volta, xml)

    def test_declaracao_original_mantida(self):
        xml = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n<a><b/></a>'
        _, volta = self.ida_e_volta(xml)
        self.assertTrue(volta.startswith(b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n<a>'))


class TestIdaEVolta(Base):
    def roundtrip(self, gerar, ext):
        orig = self.p("orig." + ext)
        gerar(orig)
        pasta = self.p("pasta")
        codigo, saida = rodar("desempacotar.py", orig, pasta)
        self.assertEqual(codigo, 0, saida)
        saida_arq = self.p("saida." + ext)
        codigo, saida = rodar("empacotar.py", pasta, saida_arq, "--original", orig)
        self.assertEqual(codigo, 0, saida)
        self.assertIn("mudadas (0)", saida)
        self.assertIn("APROVADO", saida)
        with zipfile.ZipFile(saida_arq) as zf:
            nomes = zf.namelist()
        self.assertEqual(nomes[0], "[Content_Types].xml")
        self.assertEqual(nomes[1], "_rels/.rels")
        return orig, pasta, saida_arq

    def test_docx(self):
        orig = self.p("orig.docx")
        d = docx.Document()
        d.add_paragraph("Olá mundo")
        d.add_paragraph().add_run("  espaço nas pontas  ")
        d.add_paragraph().add_run(" ")
        d.save(orig)
        pasta = self.p("pasta")
        self.assertEqual(rodar("desempacotar.py", orig, pasta)[0], 0)
        with open(os.path.join(pasta, "word", "document.xml"), "rb") as f:
            self.assertIn(b"\n  <w:body>", f.read())
        saida = self.p("saida.docx")
        codigo, txt = rodar("empacotar.py", pasta, saida, "--original", orig)
        self.assertEqual(codigo, 0, txt)
        self.assertEqual([p.text for p in docx.Document(saida).paragraphs],
                         ["Olá mundo", "  espaço nas pontas  ", " "])
        self.assertEqual(ler_parte(saida, "word/document.xml"), ler_parte(orig, "word/document.xml"))

    def test_pptx(self):
        _, _, saida = self.roundtrip(gerar_pptx, "pptx")
        titulos = [s.shapes.title.text for s in pptx.Presentation(saida).slides]
        self.assertEqual(titulos, ["Título 1", "Título 2"])

    def test_xlsx(self):
        _, _, saida = self.roundtrip(gerar_xlsx, "xlsx")
        wb = openpyxl.load_workbook(saida)
        self.assertEqual(wb["Dados"]["B1"].value, "=1+1")
        self.assertEqual(wb.sheetnames, ["Dados", "Outra"])

    def test_edicao_aparece_como_mudada(self):
        orig, pasta, _ = self.roundtrip(gerar_pptx, "pptx")
        slide = os.path.join(pasta, "ppt", "slides", "slide1.xml")
        with open(slide, "rb") as f:
            dados = f.read()
        with open(slide, "wb") as f:
            f.write(dados.replace("Título 1".encode(), "Título novo".encode()))
        codigo, saida = rodar("empacotar.py", pasta, self.p("editado.pptx"), "--original", orig)
        self.assertEqual(codigo, 0, saida)
        self.assertIn("mudadas (1): ppt/slides/slide1.xml", saida)
        self.assertEqual(pptx.Presentation(self.p("editado.pptx")).slides[0].shapes.title.text, "Título novo")


class TestDesempacotarSeguro(Base):
    def zip_malicioso(self, nome_entrada, link=False):
        caminho = self.p("mal.docx")
        with zipfile.ZipFile(caminho, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            info = zipfile.ZipInfo(nome_entrada)
            if link:
                info.create_system = 3
                info.external_attr = (0o120777) << 16
            zf.writestr(info, "conteudo")
        return caminho

    def test_zip_slip_recusado(self):
        for nome in ("../fora.txt", "a/../../fora.txt", "/abs.txt", "a\\..\\b.txt", "C:/x.txt"):
            with self.subTest(nome=nome):
                pasta = self.p("destino-" + str(abs(hash(nome))))
                codigo, saida = rodar("desempacotar.py", self.zip_malicioso(nome), pasta)
                self.assertEqual(codigo, 1, saida)
                self.assertFalse(os.path.exists(self.p("fora.txt")))
                self.assertFalse(os.path.exists(os.path.join(pasta, "[Content_Types].xml")))

    def test_link_simbolico_recusado(self):
        codigo, saida = rodar("desempacotar.py", self.zip_malicioso("word/link", link=True), self.p("d"))
        self.assertEqual(codigo, 1)
        self.assertIn("link simbólico", saida)

    def test_nome_repetido_recusado(self):
        caminho = self.p("rep.docx")
        with zipfile.ZipFile(caminho, "w") as zf:
            zf.writestr("a.xml", "<a/>")
            zf.writestr("A.xml", "<b/>")
        self.assertEqual(rodar("desempacotar.py", caminho, self.p("d"))[0], 1)

    def test_pasta_com_conteudo_e_arquivo_faltando(self):
        orig = self.p("o.docx")
        gerar_docx(orig)
        os.makedirs(self.p("cheia"))
        open(self.p("cheia", "x"), "w").close()
        self.assertEqual(rodar("desempacotar.py", orig, self.p("cheia"))[0], 2)
        self.assertEqual(rodar("desempacotar.py", self.p("nao-existe.docx"), self.p("d"))[0], 3)

    def test_xml_quebrado_gravado_como_veio(self):
        orig, ruim = self.p("o.docx"), self.p("ruim.docx")
        gerar_docx(orig)
        trocar_partes(orig, ruim, extras={"word/quebrado.xml": b"<a><b></a>"})
        codigo, saida = rodar("desempacotar.py", ruim, self.p("d"))
        self.assertEqual(codigo, 0)
        self.assertIn("quebrado.xml", saida)
        with open(self.p("d", "word", "quebrado.xml"), "rb") as f:
            self.assertEqual(f.read(), b"<a><b></a>")


class TestEmpacotar(Base):
    def preparar(self):
        orig = self.p("o.docx")
        gerar_docx(orig)
        pasta = self.p("pasta")
        self.assertEqual(rodar("desempacotar.py", orig, pasta)[0], 0)
        return orig, pasta

    def test_ignora_lixo_e_avisa(self):
        _, pasta = self.preparar()
        open(os.path.join(pasta, ".DS_Store"), "w").close()
        open(os.path.join(pasta, "word", "document.xml~"), "w").close()
        codigo, saida = rodar("empacotar.py", pasta, self.p("s.docx"))
        self.assertEqual(codigo, 0, saida)
        self.assertIn("ignorado .DS_Store", saida)
        nomes = zipfile.ZipFile(self.p("s.docx")).namelist()
        self.assertNotIn(".DS_Store", nomes)
        self.assertNotIn("word/document.xml~", nomes)

    def test_relacao_quebrada_reprova(self):
        _, pasta = self.preparar()
        os.remove(os.path.join(pasta, "word", "styles.xml"))
        codigo, saida = rodar("empacotar.py", pasta, self.p("s.docx"))
        self.assertEqual(codigo, 1)
        self.assertIn("word/styles.xml, que não existe", saida)
        self.assertTrue(os.path.exists(self.p("s.docx")))

    def test_saida_dentro_da_pasta(self):
        _, pasta = self.preparar()
        self.assertEqual(rodar("empacotar.py", pasta, os.path.join(pasta, "s.docx"))[0], 2)

    def test_recuo_acidental_avisado(self):
        _, pasta = self.preparar()
        doc = os.path.join(pasta, "word", "document.xml")
        with open(doc, encoding="utf-8") as f:
            texto = f.read()
        with open(doc, "w", encoding="utf-8") as f:
            f.write(texto.replace("<w:t>Olá mundo</w:t>", "<w:t>\n        Olá mundo\n      </w:t>"))
        codigo, saida = rodar("empacotar.py", pasta, self.p("s.docx"))
        self.assertEqual(codigo, 0, saida)
        self.assertIn("provável recuo por engano", saida)

    def test_sobrescreve_sem_sobra(self):
        orig, pasta = self.preparar()
        destino = self.p("s.docx")
        trocar_partes(orig, destino, extras={"sobra/velha.xml": b"<a/>"})
        self.assertEqual(rodar("empacotar.py", pasta, destino)[0], 0)
        self.assertNotIn("sobra/velha.xml", zipfile.ZipFile(destino).namelist())


class TestConferirEstrutura(Base):
    def test_nao_zip(self):
        with open(self.p("x.docx"), "w") as f:
            f.write("não sou zip")
        codigo, saida = rodar("conferir.py", self.p("x.docx"))
        self.assertEqual(codigo, 1)
        self.assertIn("não é um zip válido", saida)

    def test_arquivo_faltando(self):
        self.assertEqual(rodar("conferir.py", self.p("nada.docx"))[0], 3)

    def test_crc_corrompido(self):
        orig = self.p("o.docx")
        gerar_docx(orig)
        with zipfile.ZipFile(orig) as zf:
            info = zf.getinfo("word/document.xml")
            inicio = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
        with open(orig, "r+b") as f:
            f.seek(inicio + 5)
            byte = f.read(1)
            f.seek(inicio + 5)
            f.write(bytes([byte[0] ^ 0xFF]))
        codigo, saida = rodar("conferir.py", orig)
        self.assertEqual(codigo, 1, saida)

    def test_xml_malformado_e_dtd(self):
        orig, ruim = self.p("o.docx"), self.p("r.docx")
        gerar_docx(orig)
        trocar_partes(orig, ruim, trocas={
            "word/styles.xml": b"<w:styles>",
            "docProps/app.xml": b'<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x "y">]><a>&x;</a>'})
        codigo, saida = rodar("conferir.py", ruim)
        self.assertEqual(codigo, 1)
        self.assertIn("word/styles.xml: XML malformado", saida)
        self.assertIn("declaração de DTD", saida)

    def test_tipos_de_conteudo(self):
        orig, ruim = self.p("o.docx"), self.p("r.docx")
        gerar_docx(orig)

        def mexer(dados):
            return dados.replace(b"</Types>", b'<Override PartName="/word/fantasma.xml" '
                                 b'ContentType="application/xml"/></Types>')
        trocar_partes(orig, ruim, trocas={"[Content_Types].xml": mexer}, extras={"word/solto.bin": b"x"})
        codigo, saida = rodar("conferir.py", ruim)
        self.assertEqual(codigo, 1)
        self.assertIn("Override aponta para parte que não existe: /word/fantasma.xml", saida)
        self.assertIn("word/solto.bin sem tipo de conteúdo", saida)

    def test_rels_de_parte_inexistente_e_alvo_fora(self):
        orig, ruim = self.p("o.docx"), self.p("r.docx")
        gerar_docx(orig)
        rels = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
                'relationships"><Relationship Id="rId1" Type="x" Target="../../fora.xml"/></Relationships>')
        trocar_partes(orig, ruim, extras={"word/_rels/nada.xml.rels": rels})
        codigo, saida = rodar("conferir.py", ruim)
        self.assertEqual(codigo, 1)
        self.assertIn("descreve a parte word/nada.xml, que não existe", saida)
        self.assertIn("aponta para fora do pacote", saida)

    def test_externo_nao_e_cobrado(self):
        orig, ok = self.p("o.docx"), self.p("ok.docx")
        gerar_docx(orig)

        def mexer(dados):
            return dados.replace(b"</Relationships>", b'<Relationship Id="rId99" Type="http://schemas.'
                                 b'openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
                                 b'Target="https://exemplo.com/a b" TargetMode="External"/></Relationships>')
        trocar_partes(orig, ok, trocas={"word/_rels/document.xml.rels": mexer})
        self.assertEqual(rodar("conferir.py", ok)[0], 0)

    def test_pptx_ids_repetidos(self):
        orig, ruim = self.p("o.pptx"), self.p("r.pptx")
        gerar_pptx(orig)

        def slide_ids(dados):
            raiz = etree.fromstring(dados)
            ids = raiz.findall(".//{*}sldId")
            ids[1].set("id", ids[0].get("id"))
            return etree.tostring(raiz)

        def forma_ids(dados):
            raiz = etree.fromstring(dados)
            formas = raiz.findall(".//{*}cNvPr")
            formas[1].set("id", formas[0].get("id"))
            return etree.tostring(raiz)
        trocar_partes(orig, ruim, trocas={"ppt/presentation.xml": slide_ids, "ppt/slides/slide1.xml": forma_ids})
        codigo, saida = rodar("conferir.py", ruim)
        self.assertEqual(codigo, 1)
        self.assertIn("id de slide", saida)
        self.assertIn("ppt/slides/slide1.xml: id de forma", saida)

    def test_pptx_slide_sem_parte(self):
        orig, ruim = self.p("o.pptx"), self.p("r.pptx")
        gerar_pptx(orig)
        trocar_partes(orig, ruim, apagar=("ppt/slides/slide2.xml", "ppt/slides/_rels/slide2.xml.rels"))
        codigo, saida = rodar("conferir.py", ruim)
        self.assertEqual(codigo, 1)
        self.assertIn("sem parte correspondente", saida)

    def test_docx_estilo_repetido_e_comentario_orfao(self):
        orig, ruim = self.p("o.docx"), self.p("r.docx")
        gerar_docx(orig)

        def estilos(dados):
            raiz = etree.fromstring(dados)
            st = raiz.findall("{%s}style" % W)
            st[1].set("{%s}styleId" % W, st[0].get("{%s}styleId" % W))
            return etree.tostring(raiz)

        def doc(dados):
            def f(raiz):
                p = paragrafos(raiz)[0]
                for e in fragmento('<w:r><w:commentReference w:id="7"/></w:r>'):
                    p.append(e)
            return editar_documento(dados, f)
        trocar_partes(orig, ruim, trocas={"word/styles.xml": estilos, "word/document.xml": doc})
        codigo, saida = rodar("conferir.py", ruim)
        self.assertEqual(codigo, 1)
        self.assertIn("estilo", saida)
        self.assertIn("commentReference id 7 sem comentário", saida)

    def test_xlsx_aba_repetida(self):
        orig, ruim = self.p("o.xlsx"), self.p("r.xlsx")
        gerar_xlsx(orig)
        trocar_partes(orig, ruim, trocas={"xl/workbook.xml": lambda d: d.replace(b'name="Outra"', b'name="DADOS"')})
        codigo, saida = rodar("conferir.py", ruim)
        self.assertEqual(codigo, 1)
        self.assertIn("aba 'dados' repetida", saida)

    def test_problema_do_original_nao_reprova(self):
        orig_bom, orig_ruim, novo = self.p("bom.docx"), self.p("modelo.docx"), self.p("novo.docx")
        gerar_docx(orig_bom)
        trocar_partes(orig_bom, orig_ruim, extras={"word/solto.bin": b"x"})
        trocar_partes(orig_ruim, novo, trocas={"word/document.xml": lambda d: d.replace(b"mundo", b"terra")})
        codigo, saida = rodar("conferir.py", novo)
        self.assertEqual(codigo, 1)
        codigo, saida = rodar("conferir.py", novo, "--original", orig_ruim)
        self.assertEqual(codigo, 0, saida)
        self.assertIn("já existiam no original", saida)
        self.assertIn("mudadas (1): word/document.xml", saida)

    def test_importavel(self):
        orig = self.p("o.xlsx")
        gerar_xlsx(orig)
        rel = conferir.conferir(orig)
        self.assertFalse(rel.reprovou)
        self.assertEqual(rel.formato, "xlsx")


class TestConferirAutor(Base):
    def setUp(self):
        super().setUp()
        self.orig = self.p("orig.docx")
        gerar_docx(self.orig, ("Olá mundo", "Segundo parágrafo", "Terceiro"))

    def com_documento(self, funcao, nome="novo.docx"):
        destino = self.p(nome)
        trocar_partes(self.orig, destino, trocas={"word/document.xml": lambda d: editar_documento(d, funcao)})
        return destino

    def test_troca_marcada_pelo_autor_passa(self):
        def f(raiz):
            substituir_conteudo(paragrafos(raiz)[0],
                                '<w:r><w:t xml:space="preserve">Olá </w:t></w:r>'
                                '<w:del w:id="1" w:author="Agente" w:date="%s"><w:r><w:delText>mundo</w:delText></w:r></w:del>'
                                '<w:ins w:id="2" w:author="Agente" w:date="%s"><w:r><w:t>planeta</w:t></w:r></w:ins>'
                                % (DATA, DATA))
        novo = self.com_documento(f)
        codigo, saida = rodar("conferir.py", novo, "--original", self.orig, "--autor", "Agente")
        self.assertEqual(codigo, 0, saida)

    def test_autor_errado_reprova(self):
        def f(raiz):
            substituir_conteudo(paragrafos(raiz)[0],
                                '<w:r><w:t>Olá mundo</w:t></w:r>'
                                '<w:ins w:id="2" w:author="Outra Pessoa" w:date="%s"><w:r><w:t>!</w:t></w:r></w:ins>' % DATA)
        novo = self.com_documento(f)
        codigo, saida = rodar("conferir.py", novo, "--original", self.orig, "--autor", "Agente")
        self.assertEqual(codigo, 1)
        self.assertIn("do autor 'Outra Pessoa', esperado 'Agente'", saida)

    def test_mudanca_sem_marca_reprova(self):
        def f(raiz):
            substituir_conteudo(paragrafos(raiz)[1], "<w:r><w:t>Segundo parágrafo mexido</w:t></w:r>")
        novo = self.com_documento(f)
        codigo, saida = rodar("conferir.py", novo, "--original", self.orig, "--autor", "Agente")
        self.assertEqual(codigo, 1)
        self.assertIn("mudança sem marca de revisão", saida)
        self.assertIn("+Segundo parágrafo mexido", saida)

    def test_paragrafo_apagado_e_inserido_pelo_autor(self):
        def f(raiz):
            pars = paragrafos(raiz)
            alvo = pars[1]
            substituir_conteudo(alvo, '<w:del w:id="3" w:author="Agente" w:date="%s"><w:r>'
                                '<w:delText>Segundo parágrafo</w:delText></w:r></w:del>' % DATA)
            ppr = alvo.find("{%s}pPr" % W)
            if ppr is None:
                ppr = etree.SubElement(alvo, "{%s}pPr" % W)
                alvo.insert(0, ppr)
            rpr = etree.SubElement(ppr, "{%s}rPr" % W)
            rpr.append(fragmento('<w:del w:id="4" w:author="Agente" w:date="%s"/>' % DATA)[0])
            novo_p = fragmento('<w:p><w:pPr><w:rPr><w:ins w:id="5" w:author="Agente" w:date="%s"/></w:rPr></w:pPr>'
                               '<w:ins w:id="6" w:author="Agente" w:date="%s"><w:r><w:t>Parágrafo novo</w:t></w:r></w:ins></w:p>'
                               % (DATA, DATA))[0]
            pars[2].addprevious(novo_p)
        novo = self.com_documento(f)
        codigo, saida = rodar("conferir.py", novo, "--original", self.orig, "--autor", "Agente")
        self.assertEqual(codigo, 0, saida)

    def test_alteracao_previa_de_outro_autor_tolerada(self):
        def outro(raiz):
            substituir_conteudo(paragrafos(raiz)[2], '<w:r><w:t>Terceiro</w:t></w:r>'
                                '<w:ins w:id="9" w:author="Revisora" w:date="%s"><w:r><w:t> visto</w:t></w:r></w:ins>' % DATA)
        base = self.com_documento(outro, "base.docx")
        self.orig = base

        def meu(raiz):
            outro(raiz)
            substituir_conteudo(paragrafos(raiz)[0], '<w:r><w:t>Olá mundo</w:t></w:r>'
                                '<w:ins w:id="10" w:author="Agente" w:date="%s"><w:r><w:t>!</w:t></w:r></w:ins>' % DATA)
        novo = self.com_documento(meu)
        codigo, saida = rodar("conferir.py", novo, "--original", base, "--autor", "Agente")
        self.assertEqual(codigo, 0, saida)

    def test_texto_rejeitado_funcao(self):
        raiz = etree.fromstring(
            ('<w:document xmlns:w="%s"><w:body>'
             '<w:p><w:r><w:t>A</w:t></w:r><w:ins w:author="X"><w:r><w:t>B</w:t></w:r></w:ins>'
             '<w:del w:author="Y"><w:r><w:delText>C</w:delText></w:r></w:del>'
             '<w:del w:author="X"><w:r><w:delText>D</w:delText></w:r></w:del><w:r><w:tab/><w:t>E</w:t></w:r></w:p>'
             '</w:body></w:document>' % W).encode())
        self.assertEqual(conferir.texto_com_autor_rejeitado(raiz, "X"), ["AD\tE"])


@unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "LibreOffice ausente")
class TestSofficeSeguro(Base):
    def test_converter_docx_para_pdf(self):
        orig = self.p("doc.docx")
        gerar_docx(orig)
        pdf = soffice_seguro.converter(orig, "pdf", self.p("saida"), tempo_limite=180)
        self.assertTrue(pdf.is_file())
        self.assertEqual(pdf.suffix, ".pdf")
        with open(pdf, "rb") as f:
            self.assertEqual(f.read(5), b"%PDF-")
        self.assertEqual(sorted(os.listdir(self.tmp)), ["doc.docx", "saida"])

    def test_linha_de_comando(self):
        orig = self.p("deck.pptx")
        gerar_pptx(orig, 1)
        codigo, saida = rodar("soffice_seguro.py", "--converter-para", "pdf", orig, "--saida", self.p("render"))
        self.assertEqual(codigo, 0, saida)
        self.assertTrue(os.path.isfile(self.p("render", "deck.pdf")))

    def test_tempo_esgotado(self):
        orig = self.p("doc.docx")
        gerar_docx(orig)
        with self.assertRaises(soffice_seguro.ErroSoffice) as ctx:
            soffice_seguro.converter(orig, "pdf", self.p("s"), tempo_limite=0.05)
        self.assertEqual(ctx.exception.codigo, 4)

    def test_nao_grava_por_cima_da_entrada(self):
        orig = self.p("doc.docx")
        gerar_docx(orig)
        with self.assertRaises(soffice_seguro.ErroSoffice) as ctx:
            soffice_seguro.converter(orig, "docx")
        self.assertEqual(ctx.exception.codigo, 2)

    def test_entrada_que_nao_converte(self):
        ruim = self.p("lixo.docx")
        with open(ruim, "wb") as f:
            f.write(b"\x00\x01 isto nao e documento")
        with self.assertRaises(soffice_seguro.ErroSoffice) as ctx:
            soffice_seguro.converter(ruim, "pdf", self.p("s"), tempo_limite=120)
        self.assertEqual(ctx.exception.codigo, 4)


class TestSofficeAusente(Base):
    def test_sem_soffice_sai_3(self):
        orig = self.p("doc.docx")
        gerar_docx(orig)
        amb = dict(os.environ, PATH=self.tmp, SOFFICE="")
        proc = subprocess.run([sys.executable, os.path.join(SCRIPTS, "soffice_seguro.py"),
                               "--converter-para", "pdf", orig], capture_output=True, text=True, env=amb)
        self.assertEqual(proc.returncode, 3, proc.stderr)
        self.assertIn("não achei o LibreOffice", proc.stderr)

    def test_modo_obrigatorio(self):
        self.assertEqual(rodar("soffice_seguro.py")[0], 2)


if __name__ == "__main__":
    unittest.main()
