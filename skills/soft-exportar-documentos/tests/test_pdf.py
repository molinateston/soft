"""Testes da pasta scripts/pdf/.

Rodar com o Python que tem pypdf (por exemplo o venv):
    python -m unittest discover -s tests -p 'test_pdf.py'
Os PDFs de teste sao gerados na hora por tests/_fabrica_pdf.py.
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(AQUI), "scripts", "pdf")
sys.path.insert(0, AQUI)
sys.path.insert(0, SCRIPTS)

try:
    import pypdf  # noqa: F401
except ImportError:
    raise unittest.SkipTest("sem pypdf neste Python; rode com o Python do venv que tem pypdf")

import _fabrica_pdf as fab  # noqa: E402
import _comum as c  # noqa: E402
from pypdf import PdfReader, PdfWriter  # noqa: E402
from pypdf.generic import NameObject, NumberObject, RectangleObject  # noqa: E402


def rodar(script, *args, python=None):
    cmd = [python or sys.executable, os.path.join(SCRIPTS, script)] + [str(a) for a in args]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    return r.returncode, r.stdout + r.stderr


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="teste-pdf-")
        cls.form = os.path.join(cls.tmp, "form.pdf")
        cls.chapado = os.path.join(cls.tmp, "chapado.pdf")
        cls.girado = os.path.join(cls.tmp, "girado.pdf")
        fab.criar_formulario(cls.form)
        fab.criar_chapado(cls.chapado)
        fab.criar_chapado(cls.girado, giro=90)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def arq(self, nome):
        return os.path.join(self.tmp, nome)

    def json_em(self, nome, dados):
        caminho = self.arq(nome)
        with open(caminho, "w", encoding="utf-8") as f:
            json.dump(dados, f, ensure_ascii=False)
        return caminho

    def campos_de(self, pdf):
        cod, saida = rodar("campos_info.py", pdf, "-o", self.arq("info.json"))
        self.assertEqual(cod, 0, saida)
        with open(self.arq("info.json"), encoding="utf-8") as f:
            return {cp["nome"]: cp for cp in json.load(f)["campos"]}

    def qpdf_ok(self, pdf):
        if not shutil.which("qpdf"):
            return
        r = subprocess.run(["qpdf", "--check", pdf], capture_output=True, text=True)
        self.assertIn(r.returncode, (0, 3), r.stdout + r.stderr)  # 3 = so avisos


# ------------------------------------------------------------- comum

class TestComum(unittest.TestCase):
    def test_coordenadas_ida_e_volta_em_todos_os_giros(self):
        for giro in (0, 90, 180, 270):
            info = {"x0": 10.0, "y0": 20.0, "x1": 622.0, "y1": 812.0, "giro": giro}
            ret = [100.0, 200.0, 180.0, 230.0]
            topo = c.ret_pdf_para_topo(info, ret)
            volta = c.ret_topo_para_pdf(info, topo)
            for a, b in zip(ret, volta):
                self.assertAlmostEqual(a, b, places=6, msg=f"giro {giro}")
            w, h = c.tamanho_exibido(info)
            self.assertTrue(0 <= topo[0] < topo[2] <= w and 0 <= topo[1] < topo[3] <= h)

    def test_topo_sem_giro(self):
        info = {"x0": 0.0, "y0": 0.0, "x1": 612.0, "y1": 792.0, "giro": 0}
        self.assertEqual(c.ret_pdf_para_topo(info, [72, 700, 172, 720]), [72, 72, 172, 92])

    def test_encaixar_quebra_e_reduz(self):
        m = c.Medidor()
        tam, linhas, cabe = c.encaixar("uma frase com varias palavras para quebrar", 80, 60, 10, 6, m)
        self.assertTrue(cabe)
        self.assertGreater(len(linhas), 1)
        tam, linhas, cabe = c.encaixar("texto comprido demais para caixa", 60, 9, 10, 6, m)
        self.assertLess(tam, 10)
        tam, _, cabe = c.encaixar("nao cabe de jeito nenhum aqui dentro", 20, 5, 10, 6, m)
        self.assertFalse(cabe)
        self.assertEqual(tam, 6)

    def test_literal_pdf_escapa_e_codifica(self):
        lit, trocados = c.literal_pdf("a(b)\\ção")
        self.assertEqual(lit, b"(a\\(b\\)\\\\\\347\\343o)")
        self.assertEqual(trocados, [])
        _, trocados = c.literal_pdf("ok 中")
        self.assertEqual(trocados, ["中"])


class TestAjuda(unittest.TestCase):
    def test_todos_respondem_ajuda(self):
        for s in ("campos_info.py", "estrutura.py", "preencher_campos.py", "preencher_texto.py",
                  "conferir_caixas.py", "imagem_de_conferencia.py", "para_imagens.py"):
            cod, saida = rodar(s, "-h")
            self.assertEqual(cod, 0, s)
            self.assertIn("usage", saida.lower())

    def test_argumento_faltando_sai_2(self):
        cod, _ = rodar("preencher_campos.py", "so_um.pdf")
        self.assertEqual(cod, 2)

    def test_sem_pypdf_sai_3(self):
        # -S tira os site-packages, entao pypdf some
        r = subprocess.run([sys.executable, "-S", os.path.join(SCRIPTS, "campos_info.py"), "x.pdf"],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 3)
        self.assertIn("pip install pypdf", r.stderr)


# --------------------------------------------------------- campos_info

class TestCamposInfo(Base):
    def test_so_verificar_com_campos(self):
        cod, saida = rodar("campos_info.py", self.form, "--so-verificar")
        self.assertEqual(cod, 0)
        self.assertIn("tem 8 campos", saida)

    def test_so_verificar_sem_campos(self):
        cod, saida = rodar("campos_info.py", self.chapado, "--so-verificar")
        self.assertEqual(cod, 0)
        self.assertIn("nao tem campos", saida)

    def test_tipos_e_nomes_aninhados(self):
        cp = self.campos_de(self.form)
        esperado = {"nome": "texto", "endereco.rua": "texto", "endereco.cidade": "texto",
                    "aceito": "caixa_selecao", "plano": "radio", "estado": "lista_suspensa",
                    "interesses": "lista", "enviar": "botao"}
        self.assertEqual({k: v["tipo"] for k, v in cp.items()}, esperado)
        self.assertEqual(cp["endereco.rua"]["tamanho_fonte"], 10.0)  # herdado do pai

    def test_pagina_e_retangulos(self):
        cp = self.campos_de(self.form)
        self.assertEqual(cp["nome"]["pagina"], 1)
        self.assertEqual(cp["plano"]["pagina"], 2)
        self.assertEqual(cp["nome"]["retangulo_pdf"], [120.0, 695.0, 400.0, 715.0])
        self.assertEqual(cp["nome"]["retangulo_imagem"], [120.0, 77.0, 400.0, 97.0])
        self.assertEqual(len(cp["plano"]["widgets"]), 2)

    def test_dpi_escala_retangulo_imagem(self):
        cod, saida = rodar("campos_info.py", self.form, "--dpi", "144")
        self.assertEqual(cod, 0)
        cp = {x["nome"]: x for x in json.loads(saida)["campos"]}
        self.assertEqual(cp["nome"]["retangulo_imagem"], [240.0, 154.0, 800.0, 194.0])

    def test_opcoes_e_valores(self):
        cp = self.campos_de(self.form)
        self.assertEqual(cp["aceito"]["valor_ligado"], "Sim")
        self.assertEqual(cp["aceito"]["valor"], "Off")
        self.assertEqual(cp["plano"]["opcoes"], ["Basico", "Premium"])
        self.assertEqual(cp["estado"]["opcoes"], ["SP", "RJ", "MG"])
        self.assertEqual(cp["interesses"]["opcoes"], ["a", "b"])
        self.assertEqual(cp["interesses"]["textos_opcoes"], ["Alfa", "Beta"])
        self.assertTrue(cp["interesses"]["multipla_escolha"])
        self.assertEqual(cp["nome"]["tamanho_maximo"], 40)
        self.assertEqual(cp["nome"]["descricao"], "Nome completo")

    def test_arquivo_inexistente_sai_3(self):
        cod, saida = rodar("campos_info.py", self.arq("nao_existe.pdf"))
        self.assertEqual(cod, 3)
        self.assertIn("nao encontrado", saida)


# ---------------------------------------------------- preencher_campos

class TestPreencherCampos(Base):
    VALORES = {"nome": "José da Silva", "endereco.rua": "Rua (A), 10",
               "endereco.cidade": "São Paulo", "aceito": "Sim", "plano": "Premium",
               "estado": "RJ", "interesses": ["a", "Beta"]}

    def test_preenche_todos_os_tipos(self):
        saida_pdf = self.arq("preenchido.pdf")
        cod, saida = rodar("preencher_campos.py", self.form,
                           self.json_em("v.json", self.VALORES), saida_pdf)
        self.assertEqual(cod, 0, saida)
        cp = self.campos_de(saida_pdf)
        self.assertEqual(cp["nome"]["valor"], "José da Silva")
        self.assertEqual(cp["endereco.cidade"]["valor"], "São Paulo")
        self.assertEqual(cp["aceito"]["valor"], "Sim")
        self.assertEqual(cp["aceito"]["widgets"][0]["estado_atual"], "Sim")
        self.assertEqual(cp["plano"]["valor"], "Premium")
        self.assertEqual([w["estado_atual"] for w in cp["plano"]["widgets"]], ["Off", "Premium"])
        self.assertEqual(cp["estado"]["valor"], "RJ")
        self.assertEqual(cp["interesses"]["valor"], ["a", "b"])  # texto exibido vira valor
        leitor = PdfReader(saida_pdf)
        self.assertTrue(leitor.root_object["/AcroForm"]["/NeedAppearances"])
        anots = [a.get_object() for a in leitor.pages[0]["/Annots"]]
        self.assertTrue(all("/AP" in a for a in anots))
        self.qpdf_ok(saida_pdf)

    def test_aparencia_aparece_na_renderizacao(self):
        if not shutil.which("pdftotext"):
            self.skipTest("sem pdftotext")
        saida_pdf = self.arq("preenchido2.pdf")
        rodar("preencher_campos.py", self.form, self.json_em("v.json", self.VALORES), saida_pdf)
        # desliga o NeedAppearances para ver so a aparencia gerada pelo script
        w = PdfWriter(clone_from=saida_pdf)
        from pypdf.generic import BooleanObject
        w.root_object["/AcroForm"][NameObject("/NeedAppearances")] = BooleanObject(False)
        sem = self.arq("sem_need.pdf")
        w.write(sem)
        texto = subprocess.run(["pdftotext", sem, "-"], capture_output=True, text=True).stdout
        self.assertIn("José da Silva", texto)
        self.assertIn("São Paulo", texto)

    def test_campo_inexistente_sai_1_sem_gravar(self):
        saida_pdf = self.arq("nao_deve_existir.pdf")
        cod, saida = rodar("preencher_campos.py", self.form,
                           self.json_em("v.json", {"fantasma": "x", "nome": "a"}), saida_pdf)
        self.assertEqual(cod, 1)
        self.assertIn("Campo nao encontrado: fantasma", saida)
        self.assertIn("endereco.rua", saida)  # lista os que existem
        self.assertFalse(os.path.exists(saida_pdf))

    def test_caixa_usa_o_valor_ligado_do_campo(self):
        cod, saida = rodar("preencher_campos.py", self.form,
                           self.json_em("v.json", {"aceito": "Yes"}), self.arq("x.pdf"))
        self.assertEqual(cod, 1)
        self.assertIn("'Sim'", saida)
        saida_pdf = self.arq("caixa_true.pdf")
        cod, saida = rodar("preencher_campos.py", self.form,
                           self.json_em("v.json", {"aceito": True}), saida_pdf)
        self.assertEqual(cod, 0, saida)
        self.assertEqual(self.campos_de(saida_pdf)["aceito"]["valor"], "Sim")

    def test_opcoes_invalidas_recusadas(self):
        for valores, trecho in (({"plano": "Ouro"}, "Basico"), ({"estado": "BA"}, "SP"),
                                ({"interesses": ["z"]}, "opcao 'z'"),
                                ({"estado": ["SP", "RJ"]}, "uma opcao so"),
                                ({"enviar": "x"}, "botao"),
                                ({"nome": "x" * 41}, "aceita 40")):
            cod, saida = rodar("preencher_campos.py", self.form,
                               self.json_em("v.json", valores), self.arq("x.pdf"))
            self.assertEqual(cod, 1, valores)
            self.assertIn(trecho, saida)

    def test_parcial_grava_o_que_serve(self):
        saida_pdf = self.arq("parcial.pdf")
        cod, saida = rodar("preencher_campos.py", self.form,
                           self.json_em("v.json", {"nome": "Ana", "fantasma": 1}),
                           saida_pdf, "--parcial")
        self.assertEqual(cod, 1)
        self.assertEqual(self.campos_de(saida_pdf)["nome"]["valor"], "Ana")

    def test_aceita_lista_de_objetos(self):
        saida_pdf = self.arq("lista.pdf")
        valores = {"campos": [{"nome": "endereco.rua", "descricao": "rua", "pagina": 1,
                               "valor": "Rua B"}]}
        cod, saida = rodar("preencher_campos.py", self.form,
                           self.json_em("v.json", valores), saida_pdf)
        self.assertEqual(cod, 0, saida)
        self.assertIn("Ficaram sem valor", saida)
        self.assertEqual(self.campos_de(saida_pdf)["endereco.rua"]["valor"], "Rua B")

    def test_pdf_sem_campos_sai_1(self):
        cod, saida = rodar("preencher_campos.py", self.chapado,
                           self.json_em("v.json", {"a": 1}), self.arq("x.pdf"))
        self.assertEqual(cod, 1)
        self.assertIn("preencher_texto", saida)


# ----------------------------------------------------------- estrutura

class TestEstrutura(Base):
    def ler(self, pdf, *extra):
        destino = self.arq("estrutura.json")
        cod, saida = rodar("estrutura.py", pdf, "-o", destino, *extra)
        self.assertEqual(cod, 0, saida)
        with open(destino, encoding="utf-8") as f:
            return json.load(f)

    def test_rotulos_linhas_e_quadrados(self):
        d = self.ler(self.chapado)
        p = d["paginas"][0]
        self.assertEqual((p["largura"], p["altura"], p["giro"]), (612.0, 792.0, 0))
        linhas = [l for b in p["blocos"] for l in b["linhas"]]
        textos = [l["texto"] for l in linhas]
        self.assertIn("Nome completo:", textos)
        nome = linhas[textos.index("Nome completo:")]
        self.assertAlmostEqual(nome["caixa"][0], 72.0, delta=0.5)
        self.assertTrue(80 < nome["caixa"][1] < nome["caixa"][3] < 96)
        self.assertEqual([w["texto"] for w in nome["palavras"]], ["Nome", "completo:"])
        self.assertEqual(len(p["linhas_horizontais"]), 2)
        self.assertEqual(p["linhas_horizontais"][0]["caixa"], [170.0, 94.0, 540.0, 94.0])
        self.assertEqual(p["quadrados"], [{"caixa": [180.0, 164.0, 190.0, 174.0]}])
        self.assertEqual(len(d["paginas"]), 2)

    def test_pagina_girada(self):
        p = self.ler(self.girado)["paginas"][0]
        self.assertEqual((p["largura"], p["altura"], p["giro"]), (792.0, 612.0, 90))
        self.assertEqual(len(p["linhas_verticais"]), 2)
        palavra = p["blocos"][0]["linhas"][0]
        self.assertTrue(palavra["caixa"][2] <= 792 and palavra["caixa"][3] <= 612)

    def test_linha_dentro_de_formulario_interno(self):
        from pypdf.generic import (ArrayObject, DecodedStreamObject, DictionaryObject,
                                   FloatObject)
        w = PdfWriter()
        p = w.add_blank_page(612, 792)
        xo = DecodedStreamObject()
        xo.set_data(b"0 0 200 1 re f")
        xo.update({NameObject("/Type"): NameObject("/XObject"),
                   NameObject("/Subtype"): NameObject("/Form"),
                   NameObject("/BBox"): ArrayObject([FloatObject(v) for v in (0, 0, 200, 1)])})
        conteudo = DecodedStreamObject()
        conteudo.set_data(b"q 1 0 0 1 100 500 cm /Fx Do Q")
        p[NameObject("/Contents")] = w._add_object(conteudo)
        p[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject(
            {NameObject("/Fx"): w._add_object(xo)})})
        pdf = self.arq("xobj.pdf")
        w.write(pdf)
        p = self.ler(pdf)["paginas"][0]
        self.assertEqual(p["linhas_horizontais"], [{"caixa": [100.0, 291.0, 300.0, 292.0]}])

    def test_sem_palavras_e_sem_desenhos(self):
        p = self.ler(self.chapado, "--sem-palavras", "--sem-desenhos")["paginas"][0]
        self.assertNotIn("palavras", p["blocos"][0]["linhas"][0])
        self.assertNotIn("linhas_horizontais", p)

    def test_pdf_do_libreoffice(self):
        if not shutil.which("soffice"):
            self.skipTest("sem LibreOffice")
        try:
            import docx
        except ImportError:
            self.skipTest("sem python-docx")
        doc = docx.Document()
        doc.add_paragraph("Nome do cliente: ______________________")
        tabela = doc.add_table(rows=2, cols=2)
        tabela.style = "Table Grid"
        tabela.cell(0, 0).text = "CPF"
        tabela.cell(1, 0).text = "Telefone"
        fonte = self.arq("lo.docx")
        doc.save(fonte)
        perfil = tempfile.mkdtemp(prefix="perfil-lo-", dir=self.tmp)
        subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{perfil}",
                        "--convert-to", "pdf", "--outdir", self.tmp, fonte],
                       capture_output=True, timeout=180)
        pdf = self.arq("lo.pdf")
        if not os.path.exists(pdf):
            self.skipTest("LibreOffice nao converteu")
        p = self.ler(pdf)["paginas"][0]
        textos = " ".join(l["texto"] for b in p["blocos"] for l in b["linhas"])
        self.assertIn("Nome do cliente:", textos)
        self.assertIn("Telefone", textos)
        bordas = p["linhas_horizontais"] + p["linhas_verticais"] + p["retangulos"]
        self.assertGreater(len(bordas), 0)  # bordas da tabela


# ------------------------------------------------------ preencher_texto

class TestPreencherTexto(Base):
    CAMPOS = {"sistema": "topo", "campos": [
        {"nome": "nome", "pagina": 1, "retangulo": [170, 80, 540, 94],
         "texto": "Joana Conceição (teste)", "tamanho_fonte": 11,
         "retangulo_rotulo": [72, 83.38, 159.36, 94.48]},
        {"nome": "aceite", "pagina": 1, "retangulo": [180, 162, 192, 176], "texto": "X",
         "tamanho_fonte": 9, "alinhamento": "centro", "cor": "#cc0000"},
        {"nome": "obs", "pagina": 1, "retangulo": [72, 200, 300, 240],
         "texto": "Texto longo que precisa quebrar em mais de uma linha dentro da caixa."},
        {"nome": "deixado", "pagina": 1, "retangulo": [72, 300, 200, 314], "texto": ""},
    ]}

    def anotacoes(self, pdf, pagina=0):
        p = PdfReader(pdf).pages[pagina]
        return [a.get_object() for a in p.get("/Annots", [])]

    def test_escreve_texto_livre(self):
        saida_pdf = self.arq("texto.pdf")
        cod, saida = rodar("preencher_texto.py", self.chapado,
                           self.json_em("t.json", self.CAMPOS), saida_pdf)
        self.assertEqual(cod, 0, saida)
        self.assertIn("3 textos escritos", saida)
        self.assertIn("Ficaram vazios", saida)
        anots = self.anotacoes(saida_pdf)
        self.assertEqual(len(anots), 3)
        for a in anots:
            self.assertEqual(a["/Subtype"], "/FreeText")
            self.assertEqual(a["/BS"]["/W"], 0)
            self.assertIn("/AP", a)
        rect = [float(v) for v in anots[0]["/Rect"]]
        self.assertEqual(rect, [170.0, 698.0, 540.0, 712.0])
        self.assertEqual(anots[0]["/Contents"], "Joana Conceição (teste)")
        texto = subprocess.run(["pdftotext", saida_pdf, "-"], capture_output=True,
                               text=True).stdout
        self.assertIn("Joana Conceição (teste)", texto)
        self.qpdf_ok(saida_pdf)

    def test_texto_que_nao_cabe(self):
        dados = {"campos": [{"nome": "apertado", "pagina": 1, "retangulo": [72, 260, 120, 268],
                             "texto": "Isto nao cabe de jeito nenhum aqui"}]}
        cod, saida = rodar("preencher_texto.py", self.chapado,
                           self.json_em("t.json", dados), self.arq("t.pdf"))
        self.assertEqual(cod, 1)
        self.assertIn("Texto nao coube: apertado", saida)
        cod, saida = rodar("preencher_texto.py", self.chapado,
                           self.json_em("t.json", dados), self.arq("t.pdf"), "--sem-reduzir")
        self.assertIn("fonte 10", saida)

    def test_reduz_fonte_quando_precisa(self):
        dados = {"campos": [{"nome": "curto", "pagina": 1, "retangulo": [72, 260, 140, 274],
                             "texto": "Texto mais largo", "tamanho_fonte": 12}]}
        cod, saida = rodar("preencher_texto.py", self.chapado,
                           self.json_em("t.json", dados), self.arq("t.pdf"))
        self.assertEqual(cod, 0, saida)
        self.assertIn("fonte reduzida de 12", saida)

    def test_sistema_pdf_e_escala_dpi(self):
        topo = {"campos": [{"pagina": 1, "retangulo": [100, 100, 200, 120], "texto": "a"}]}
        pdf = {"sistema": "pdf", "campos": [{"pagina": 1, "retangulo": [100, 672, 200, 692],
                                             "texto": "a"}]}
        pixels = {"escala_dpi": 144, "campos": [{"pagina": 1, "retangulo": [200, 200, 400, 240],
                                                 "texto": "a"}]}
        rects = []
        for i, dados in enumerate((topo, pdf, pixels)):
            saida_pdf = self.arq(f"s{i}.pdf")
            cod, saida = rodar("preencher_texto.py", self.chapado,
                               self.json_em("t.json", dados), saida_pdf)
            self.assertEqual(cod, 0, saida)
            rects.append([float(v) for v in self.anotacoes(saida_pdf)[0]["/Rect"]])
        self.assertEqual(rects[0], [100.0, 672.0, 200.0, 692.0])
        self.assertEqual(rects[0], rects[1])
        self.assertEqual(rects[0], rects[2])

    def test_pagina_girada(self):
        dados = {"campos": [{"pagina": 1, "retangulo": [100, 50, 300, 70], "texto": "em pe"}]}
        saida_pdf = self.arq("girado_texto.pdf")
        cod, saida = rodar("preencher_texto.py", self.girado,
                           self.json_em("t.json", dados), saida_pdf)
        self.assertEqual(cod, 0, saida)
        a = self.anotacoes(saida_pdf)[0]
        # giro 90: x do PDF = topo, y do PDF = esquerda
        self.assertEqual([float(v) for v in a["/Rect"]], [50.0, 100.0, 70.0, 300.0])
        self.assertEqual([int(v) for v in a["/AP"]["/N"]["/Matrix"]], [0, 1, -1, 0, 0, 0])
        if shutil.which("pdftotext"):
            r = subprocess.run(["pdftotext", "-bbox", saida_pdf, "-"], capture_output=True,
                               text=True).stdout
            self.assertIn(">em</word>", r)

    def test_item_invalido(self):
        dados = {"campos": [{"pagina": 9, "retangulo": [1, 1, 50, 20], "texto": "x"},
                            {"pagina": 1, "retangulo": [1, 1, 50, 20], "texto": "y",
                             "alinhamento": "meio"}]}
        cod, saida = rodar("preencher_texto.py", self.chapado,
                           self.json_em("t.json", dados), self.arq("t.pdf"))
        self.assertEqual(cod, 1)
        self.assertIn("pagina 9 nao existe", saida)
        self.assertIn("alinhamento 'meio'", saida)

    def test_json_sem_retangulo_sai_2(self):
        cod, saida = rodar("preencher_texto.py", self.chapado,
                           self.json_em("t.json", {"campos": [{"pagina": 1, "texto": "x"}]}),
                           self.arq("t.pdf"))
        self.assertEqual(cod, 2)
        self.assertIn("retangulo", saida)


# ------------------------------------------------------ conferir_caixas

class TestConferirCaixas(Base):
    def conferir(self, dados):
        return rodar("conferir_caixas.py", self.json_em("cx.json", dados))

    def test_sem_problema(self):
        cod, saida = self.conferir(TestPreencherTexto.CAMPOS)
        self.assertEqual(cod, 0, saida)

    def test_sobreposicao_entre_caixas(self):
        cod, saida = self.conferir({"campos": [
            {"nome": "a", "pagina": 1, "retangulo": [10, 10, 100, 30]},
            {"nome": "b", "pagina": 1, "retangulo": [90, 20, 200, 40]}]})
        self.assertEqual(cod, 1)
        self.assertIn("a se sobrepoe a b", saida)

    def test_caixa_contra_rotulo(self):
        cod, saida = self.conferir({"campos": [
            {"nome": "nome", "pagina": 1, "retangulo": [60, 80, 300, 96],
             "retangulo_rotulo": [10, 82, 70, 94]}]})
        self.assertEqual(cod, 1)
        self.assertIn("nome (rotulo)", saida)

    def test_encostar_nao_e_sobrepor(self):
        cod, saida = self.conferir({"campos": [
            {"nome": "a", "pagina": 1, "retangulo": [10, 10, 100, 30]},
            {"nome": "b", "pagina": 1, "retangulo": [100, 10, 200, 30]},
            {"nome": "c", "pagina": 2, "retangulo": [10, 10, 100, 30]}]})
        self.assertEqual(cod, 0, saida)

    def test_caixa_baixa_para_a_fonte(self):
        cod, saida = self.conferir({"campos": [
            {"nome": "baixa", "pagina": 1, "retangulo": [10, 10, 100, 20], "tamanho_fonte": 12}]})
        self.assertEqual(cod, 1)
        self.assertIn("baixa demais para fonte 12", saida)

    def test_aviso_de_texto_largo_nao_reprova(self):
        cod, saida = self.conferir({"campos": [
            {"nome": "x", "pagina": 1, "retangulo": [10, 10, 40, 24], "texto": "texto bem comprido"}]})
        self.assertEqual(cod, 0)
        self.assertIn("Aviso", saida)

    def test_saida_do_campos_info(self):
        self.campos_de(self.form)
        cod, saida = rodar("conferir_caixas.py", self.arq("info.json"), "--json")
        self.assertEqual(cod, 0, saida)
        self.assertEqual(json.loads(saida)["caixas"], 9)

    def test_formato_errado_sai_2(self):
        cod, _ = self.conferir({"outra": 1})
        self.assertEqual(cod, 2)


# ------------------------------------------------- imagem_de_conferencia

class TestImagemDeConferencia(Base):
    def test_desenha_caixas(self):
        from PIL import Image
        png = self.arq("conf.png")
        dados = {"campos": [{"nome": "nome", "pagina": 1, "retangulo": [170, 80, 540, 94],
                             "retangulo_rotulo": [72, 83, 160, 95]}]}
        cod, saida = rodar("imagem_de_conferencia.py", self.chapado,
                           self.json_em("cx.json", dados), "--pagina", 1, "-o", png)
        self.assertEqual(cod, 0, saida)
        self.assertIn("2 caixas desenhadas", saida)
        img = Image.open(png).convert("RGB")
        self.assertEqual(img.size, (850, 1100))
        # borda azul da caixa de texto, na escala de 100 dpi
        x, y = int(300 * 100 / 72), int(80 * 100 / 72)
        r, g, b = img.getpixel((x, y))
        self.assertTrue(b > 150 and r < 100, (r, g, b))

    def test_campos_info_em_pagina_2(self):
        self.campos_de(self.form)
        png = self.arq("conf2.png")
        cod, saida = rodar("imagem_de_conferencia.py", self.form, self.arq("info.json"),
                           "--pagina", 2, "-o", png, "--dpi", 72)
        self.assertEqual(cod, 0, saida)
        self.assertIn("5 caixas desenhadas", saida)

    def test_pagina_inexistente_sai_2(self):
        cod, _ = rodar("imagem_de_conferencia.py", self.chapado,
                       self.json_em("cx.json", {"campos": []}), "--pagina", 5,
                       "-o", self.arq("x.png"))
        self.assertEqual(cod, 2)


# --------------------------------------------------------- para_imagens

class TestParaImagens(Base):
    def test_uma_imagem_por_pagina(self):
        from PIL import Image
        pasta = self.arq("imgs")
        cod, saida = rodar("para_imagens.py", self.form, pasta, "--dpi", 50)
        self.assertEqual(cod, 0, saida)
        nomes = sorted(os.listdir(pasta))
        self.assertEqual(nomes, ["pagina-1.png", "pagina-2.png"])
        with Image.open(os.path.join(pasta, nomes[0])) as img:
            self.assertEqual(img.size, (425, 550))

    def test_nomes_ordenados_com_zeros(self):
        w = PdfWriter()
        for _ in range(12):
            w.add_blank_page(72, 72)
        pdf = self.arq("doze.pdf")
        w.write(pdf)
        pasta = self.arq("doze")
        cod, saida = rodar("para_imagens.py", pdf, pasta, "--dpi", 10, "--prefixo", "p")
        self.assertEqual(cod, 0, saida)
        nomes = sorted(os.listdir(pasta))
        self.assertEqual(len(nomes), 12)
        self.assertEqual(nomes[0], "p-01.png")
        self.assertEqual(nomes[-1], "p-12.png")

    def test_arquivo_inexistente_sai_3(self):
        cod, _ = rodar("para_imagens.py", self.arq("nada.pdf"), self.arq("y"))
        self.assertEqual(cod, 3)


if __name__ == "__main__":
    unittest.main()
