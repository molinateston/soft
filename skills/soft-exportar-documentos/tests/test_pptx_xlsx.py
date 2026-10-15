"""Testes dos scripts pptx/ e xlsx/ do Soft Office.

Rodar a partir da pasta novo/:
    python3 -m unittest tests.test_pptx_xlsx -v

Os arquivos de teste são gerados aqui mesmo (python-pptx, openpyxl, Pillow).
Os testes que precisam do LibreOffice ou do pdftoppm são pulados se eles faltarem.
"""
import hashlib
import io
import json
import os
import posixpath
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from urllib.parse import unquote

from lxml import etree

RAIZ = Path(__file__).resolve().parent.parent
SCRIPTS = RAIZ / "scripts"
MOSAICO = SCRIPTS / "pptx" / "mosaico.py"
LIMPAR = SCRIPTS / "pptx" / "limpar.py"
NOVO_SLIDE = SCRIPTS / "pptx" / "novo_slide.py"
RECALCULAR = SCRIPTS / "xlsx" / "recalcular.py"

TEM_SOFFICE = bool(shutil.which("soffice") or shutil.which("libreoffice"))
TEM_PDFTOPPM = bool(shutil.which("pdftoppm"))

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


# ---------------------------------------------------------------- utilidades

def rodar(script, *args, env=None, timeout=300):
    return subprocess.run([sys.executable, str(script), *map(str, args)], capture_output=True,
                          text=True, env=env, timeout=timeout)


def desempacotar(arquivo, pasta):
    with zipfile.ZipFile(arquivo) as z:
        z.extractall(pasta)


def empacotar(pasta, saida):
    pasta = Path(pasta)
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(pasta / "[Content_Types].xml", "[Content_Types].xml")
        for f in sorted(pasta.rglob("*")):
            if f.is_file() and f.name != "[Content_Types].xml":
                z.write(f, f.relative_to(pasta).as_posix())


def foto(caminho, cor):
    from PIL import Image
    Image.new("RGB", (64, 48), cor).save(caminho)


def deck(caminho, n=4, pasta_tmp=None, notas=True, imagem=True, grafico=True, link=True):
    """Deck de teste: slide 2 com notas, slide 3 com imagem só dele, slide 4 com gráfico,
    slide 1 com hiperlink para o slide 3."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    p = Presentation()
    for i in range(n):
        s = p.slides.add_slide(p.slide_layouts[1 if i else 0])
        s.shapes.title.text = f"Titulo {i + 1}"
    if notas and n >= 2:
        p.slides[1].notes_slide.notes_text_frame.text = "Nota do slide dois"
    if imagem and n >= 3:
        png = Path(pasta_tmp) / "foto.png"
        foto(png, (200, 30, 30))
        p.slides[2].shapes.add_picture(str(png), Inches(1), Inches(2))
    if grafico and n >= 4:
        dados = CategoryChartData()
        dados.categories = ["A", "B", "C"]
        dados.add_series("Serie", (1, 2, 3))
        p.slides[3].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2),
                                     Inches(4), Inches(3), dados)
    if link and n >= 3:
        caixa = p.slides[0].shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1))
        caixa.text = "ir para o 3"
        caixa.click_action.target_slide = p.slides[2]
    p.save(caminho)


def ler(pasta, parte):
    return etree.parse(str(Path(pasta) / parte))


def gravar(arvore, pasta, parte):
    arvore.write(str(Path(pasta) / parte), xml_declaration=True, encoding="UTF-8", standalone=True)


def tirar_da_lista(pasta, posicao):
    """Tira a entrada de número `posicao` (1 = primeira) de p:sldIdLst. Devolve (id, rId)."""
    arv = ler(pasta, "ppt/presentation.xml")
    lista = arv.getroot().find(f"{{{NS_P}}}sldIdLst")
    sld = lista[posicao - 1]
    lista.remove(sld)
    gravar(arv, pasta, "ppt/presentation.xml")
    return sld.get("id"), sld.get(f"{{{NS_R}}}id")


def instantaneo(pasta):
    pasta = Path(pasta)
    return {f.relative_to(pasta).as_posix(): hashlib.sha256(f.read_bytes()).hexdigest()
            for f in pasta.rglob("*") if f.is_file()}


def problemas_do_pacote(pasta):
    """Conferência mínima feita aqui no teste: relação interna aponta para arquivo que existe,
    toda parte tem tipo de conteúdo, ids de slide únicos."""
    pasta = Path(pasta)
    erros = []
    tipos = ler(pasta, "[Content_Types].xml").getroot()
    defaults = {d.get("Extension").lower() for d in tipos.findall(f"{{{NS_CT}}}Default")}
    overrides = {unquote(o.get("PartName")).lstrip("/").lower() for o in tipos.findall(f"{{{NS_CT}}}Override")}
    existentes = {f.relative_to(pasta).as_posix().lower() for f in pasta.rglob("*") if f.is_file()}
    for o in overrides:
        if o not in existentes:
            erros.append(f"override sem parte: {o}")
    for f in pasta.rglob("*"):
        if not f.is_file():
            continue
        rel = f.relative_to(pasta).as_posix()
        if rel == "[Content_Types].xml":
            continue
        ext = "rels" if rel.endswith(".rels") else posixpath.splitext(rel)[1].lstrip(".").lower()
        if rel.lower() not in overrides and ext not in defaults:
            erros.append(f"sem tipo: {rel}")
        if rel.endswith(".rels"):
            d, nome = posixpath.split(rel)
            dona = posixpath.join(posixpath.dirname(d), nome[:-5]) if nome != ".rels" else ""
            for r in ler(pasta, rel).getroot():
                if r.get("TargetMode") == "External":
                    continue
                t = unquote(r.get("Target"))
                alvo = t.lstrip("/") if t.startswith("/") else posixpath.normpath(
                    posixpath.join(posixpath.dirname(dona), t))
                if not (pasta / alvo).is_file():
                    erros.append(f"relação quebrada: {rel} {r.get('Id')} -> {alvo}")
    ids = [s.get("id") for s in ler(pasta, "ppt/presentation.xml").getroot().iter(f"{{{NS_P}}}sldId")]
    if len(ids) != len(set(ids)):
        erros.append(f"ids de slide repetidos: {ids}")
    return erros


def converter_pdf(arquivo, pasta):
    """Conversão própria do teste, com perfil isolado. Devolve o caminho do PDF ou None."""
    perfil = Path(pasta) / "perfil-teste"
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    r = subprocess.run([soffice, f"-env:UserInstallation={perfil.as_uri()}", "--headless", "--norestore",
                        "--convert-to", "pdf", "--outdir", str(pasta), str(arquivo)],
                       capture_output=True, timeout=180,
                       env=dict(os.environ, HOME=str(pasta), SAL_USE_VCLPLUGIN="svp"))
    pdf = Path(pasta) / (Path(arquivo).stem + ".pdf")
    return pdf if r.returncode == 0 and pdf.is_file() else None


def paginas_pdf(pdf):
    with open(pdf, "rb") as f:
        dados = f.read()
    import re
    return len(re.findall(rb"/Type\s*/Page[^s]", dados))


def path_sem(*programas_permitidos):
    """Cria uma pasta de PATH só com os programas pedidos (links para os reais)."""
    d = Path(tempfile.mkdtemp(prefix="path-"))
    for prog in programas_permitidos:
        real = shutil.which(prog)
        if real:
            os.symlink(real, d / prog)
    return d


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="teste-pptx-xlsx-"))

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


# ---------------------------------------------------------------- mosaico

@unittest.skipUnless(TEM_SOFFICE and TEM_PDFTOPPM, "precisa de soffice e pdftoppm")
class TesteMosaico(Base):
    def test_uma_imagem_ate_12_slides(self):
        from PIL import Image
        arq = self.tmp / "d.pptx"
        deck(arq, n=5, pasta_tmp=self.tmp)
        r = rodar(MOSAICO, arq, self.tmp / "saida" / "m")
        self.assertEqual(r.returncode, 0, r.stderr)
        img = self.tmp / "saida" / "m.jpg"
        self.assertTrue(img.is_file())
        self.assertIn("5 slides em 1 imagem", r.stdout)
        with Image.open(img) as im:
            self.assertEqual(im.format, "JPEG")
            self.assertGreater(im.width, 3 * 400)
            # faixa do número escura no canto de cima de cada célula
            self.assertLess(im.convert("L").getpixel((im.width // 6, 20)), 80)

    def test_varias_imagens_acima_de_12(self):
        arq = self.tmp / "d.pptx"
        deck(arq, n=14, pasta_tmp=self.tmp)
        r = rodar(MOSAICO, arq, self.tmp / "m", "--colunas", "4", "--largura", "200")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertTrue((self.tmp / "m-1.jpg").is_file())
        self.assertTrue((self.tmp / "m-2.jpg").is_file())
        self.assertFalse((self.tmp / "m.jpg").exists())
        self.assertIn("slides 13 a 14", r.stdout)
        self.assertIn("Dica", r.stdout)

    def test_slide_oculto_mantem_numeracao(self):
        from pptx import Presentation
        arq = self.tmp / "d.pptx"
        deck(arq, n=3, pasta_tmp=self.tmp)
        p = Presentation(arq)
        p.slides[1]._element.set("show", "0")
        p.save(arq)
        r = rodar(MOSAICO, arq, self.tmp / "m")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("3 slides", r.stdout)
        self.assertIn("oculto", r.stdout)

    def test_falta_soffice_sai_3(self):
        arq = self.tmp / "d.pptx"
        deck(arq, n=2, pasta_tmp=self.tmp)
        vazio = path_sem()
        r = rodar(MOSAICO, arq, self.tmp / "m", env=dict(os.environ, PATH=str(vazio)))
        self.assertEqual(r.returncode, 3)
        self.assertIn("soffice", r.stderr)

    def test_falta_pdftoppm_sai_3(self):
        arq = self.tmp / "d.pptx"
        deck(arq, n=2, pasta_tmp=self.tmp)
        so_soffice = path_sem("soffice")
        r = rodar(MOSAICO, arq, self.tmp / "m", env=dict(os.environ, PATH=str(so_soffice)))
        self.assertEqual(r.returncode, 3)
        self.assertIn("pdftoppm", r.stderr)

    def test_arquivo_inexistente_sai_3(self):
        r = rodar(MOSAICO, self.tmp / "nao.pptx", self.tmp / "m")
        self.assertEqual(r.returncode, 3)

    def test_argumento_errado_sai_2(self):
        arq = self.tmp / "d.pptx"
        deck(arq, n=2, pasta_tmp=self.tmp)
        r = rodar(MOSAICO, arq, self.tmp / "m", "--colunas", "0")
        self.assertEqual(r.returncode, 2)


# ---------------------------------------------------------------- limpar

class TesteLimpar(Base):
    def preparar(self, n=4):
        arq = self.tmp / "d.pptx"
        deck(arq, n=n, pasta_tmp=self.tmp)
        pasta = self.tmp / "d"
        desempacotar(arq, pasta)
        return pasta

    def test_remove_slide_fora_da_lista_com_midia_e_hiperlink(self):
        from pptx import Presentation
        pasta = self.preparar()
        self.assertTrue(list((pasta / "ppt" / "media").glob("*.png")))
        id_tirado, _ = tirar_da_lista(pasta, 3)  # slide 3: a imagem e o alvo do hiperlink
        # seção que ainda cita o slide tirado
        arv = ler(pasta, "ppt/presentation.xml")
        raiz = arv.getroot()
        ext_lst = raiz.find(f"{{{NS_P}}}extLst")
        if ext_lst is None:
            ext_lst = etree.SubElement(raiz, f"{{{NS_P}}}extLst")
        ext = etree.SubElement(ext_lst, f"{{{NS_P}}}ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
        sec_lst = etree.SubElement(ext, f"{{{NS_P14}}}sectionLst", nsmap={"p14": NS_P14})
        sec = etree.SubElement(sec_lst, f"{{{NS_P14}}}section", name="Tudo",
                               id="{11111111-2222-3333-4444-555555555555}")
        ids = etree.SubElement(sec, f"{{{NS_P14}}}sldIdLst")
        for s in raiz.iter(f"{{{NS_P}}}sldId"):
            etree.SubElement(ids, f"{{{NS_P14}}}sldId", id=s.get("id"))
        etree.SubElement(ids, f"{{{NS_P14}}}sldId", id=id_tirado)
        gravar(arv, pasta, "ppt/presentation.xml")

        r = rodar(LIMPAR, pasta)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("ppt/slides/slide3.xml", r.stdout)
        self.assertIn("mídia sem uso", r.stdout)
        self.assertIn("hiperlink removido", r.stdout)
        self.assertIn("entrada de seção removida", r.stdout)
        self.assertFalse((pasta / "ppt/slides/slide3.xml").exists())
        self.assertFalse((pasta / "ppt/slides/_rels/slide3.xml.rels").exists())
        self.assertFalse(list((pasta / "ppt" / "media").glob("*.png")))
        self.assertNotIn("slide3.xml", (pasta / "[Content_Types].xml").read_text())
        self.assertNotIn("slide3.xml", (pasta / "ppt/_rels/presentation.xml.rels").read_text())
        self.assertNotIn(f'id="{id_tirado}"', (pasta / "ppt/presentation.xml").read_text())
        self.assertEqual(problemas_do_pacote(pasta), [])
        saida = self.tmp / "limpo.pptx"
        empacotar(pasta, saida)
        p = Presentation(saida)
        self.assertEqual([s.shapes.title.text for s in p.slides], ["Titulo 1", "Titulo 2", "Titulo 4"])
        if TEM_SOFFICE:
            pdf = converter_pdf(saida, self.tmp)
            self.assertIsNotNone(pdf, "LibreOffice não converteu o deck limpo")
            self.assertEqual(paginas_pdf(pdf), 3)

    def test_notas_do_slide_tirado_saem(self):
        pasta = self.preparar()
        notas_antes = list((pasta / "ppt/notesSlides").glob("*.xml"))
        self.assertEqual(len(notas_antes), 1)
        tirar_da_lista(pasta, 2)
        r = rodar(LIMPAR, pasta)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("notas sem slide", r.stdout)
        self.assertFalse(list(pasta.glob("ppt/notesSlides/*.xml")))
        self.assertTrue(list(pasta.glob("ppt/notesMasters/*.xml")), "o master de notas fica")
        self.assertEqual(problemas_do_pacote(pasta), [])

    def test_grafico_do_slide_tirado_sai(self):
        pasta = self.preparar()
        self.assertTrue(list(pasta.glob("ppt/charts/*.xml")))
        tirar_da_lista(pasta, 4)
        r = rodar(LIMPAR, pasta)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertFalse(list(pasta.glob("ppt/charts/*.xml")))
        self.assertFalse(list(pasta.glob("ppt/embeddings/*")))
        self.assertEqual(problemas_do_pacote(pasta), [])

    def test_midia_orfa_e_relacao_orfa(self):
        pasta = self.preparar()
        foto(pasta / "ppt/media/sobra.png", (0, 0, 255))
        arv = ler(pasta, "ppt/slides/_rels/slide1.xml.rels")
        etree.SubElement(arv.getroot(), f"{{{NS_REL}}}Relationship", Id="rId99",
                         Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                         Target="../media/nao_existe.png")
        gravar(arv, pasta, "ppt/slides/_rels/slide1.xml.rels")
        r = rodar(LIMPAR, pasta)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertFalse((pasta / "ppt/media/sobra.png").exists())
        self.assertIn("rId99", r.stdout)
        self.assertNotIn("rId99", (pasta / "ppt/slides/_rels/slide1.xml.rels").read_text())
        self.assertEqual(problemas_do_pacote(pasta), [])

    def test_layouts_do_master_ficam(self):
        pasta = self.preparar()
        layouts = sorted(p.name for p in pasta.glob("ppt/slideLayouts/*.xml"))
        tirar_da_lista(pasta, 2)
        rodar(LIMPAR, pasta)
        self.assertEqual(sorted(p.name for p in pasta.glob("ppt/slideLayouts/*.xml")), layouts)
        self.assertTrue(list(pasta.glob("ppt/slideMasters/*.xml")))

    def test_idempotente(self):
        pasta = self.preparar()
        tirar_da_lista(pasta, 3)
        primeira = rodar(LIMPAR, pasta)
        self.assertEqual(primeira.returncode, 0, primeira.stderr)
        antes = instantaneo(pasta)
        segunda = rodar(LIMPAR, pasta)
        self.assertEqual(segunda.returncode, 0, segunda.stderr)
        self.assertIn("Nada a remover", segunda.stdout)
        self.assertEqual(instantaneo(pasta), antes)

    def test_deck_limpo_nao_muda(self):
        pasta = self.preparar()
        antes = instantaneo(pasta)
        r = rodar(LIMPAR, pasta)
        self.assertIn("Nada a remover", r.stdout)
        self.assertEqual(instantaneo(pasta), antes)

    def test_simular_nao_grava(self):
        pasta = self.preparar()
        tirar_da_lista(pasta, 3)
        antes = instantaneo(pasta)
        r = rodar(LIMPAR, pasta, "--simular")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("slide3.xml", r.stdout)
        self.assertEqual(instantaneo(pasta), antes)

    def test_pasta_errada_sai_3(self):
        (self.tmp / "vazia").mkdir()
        self.assertEqual(rodar(LIMPAR, self.tmp / "vazia").returncode, 3)
        self.assertEqual(rodar(LIMPAR, self.tmp / "nao_existe").returncode, 3)


# ---------------------------------------------------------------- novo_slide

class TesteNovoSlide(Base):
    def preparar(self, n=4):
        arq = self.tmp / "d.pptx"
        deck(arq, n=n, pasta_tmp=self.tmp)
        pasta = self.tmp / "d"
        desempacotar(arq, pasta)
        return pasta

    def abrir(self, pasta, nome="saida.pptx"):
        from pptx import Presentation
        saida = self.tmp / nome
        empacotar(pasta, saida)
        return saida, Presentation(saida)

    def test_duplica_slide_com_notas(self):
        pasta = self.preparar()
        r = rodar(NOVO_SLIDE, pasta, "--de", "slide2")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("ppt/slides/slide5.xml", r.stdout)
        self.assertIn('<p:sldId id="', r.stdout)
        self.assertTrue((pasta / "ppt/notesSlides/notesSlide2.xml").is_file())
        notas_rels = (pasta / "ppt/notesSlides/_rels/notesSlide2.xml.rels").read_text()
        self.assertIn("slide5.xml", notas_rels)
        self.assertEqual(problemas_do_pacote(pasta), [])
        saida, p = self.abrir(pasta)
        self.assertEqual(len(p.slides), 5)
        self.assertEqual(p.slides[4].shapes.title.text, "Titulo 2")
        self.assertEqual(p.slides[4].notes_slide.notes_text_frame.text, "Nota do slide dois")
        # as notas do original continuam ligadas ao original
        self.assertEqual(p.slides[1].notes_slide.notes_text_frame.text, "Nota do slide dois")
        if TEM_SOFFICE:
            pdf = converter_pdf(saida, self.tmp)
            self.assertIsNotNone(pdf, "LibreOffice não converteu o deck com slide duplicado")
            self.assertEqual(paginas_pdf(pdf), 5)

    def test_duplica_slide_com_grafico_e_imagem(self):
        pasta = self.preparar()
        graficos = len(list(pasta.glob("ppt/charts/chart*.xml")))
        midia = sorted(p.name for p in pasta.glob("ppt/media/*"))
        r = rodar(NOVO_SLIDE, pasta, "--de", "4")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(len(list(pasta.glob("ppt/charts/chart*.xml"))), graficos + 1)
        self.assertEqual(sorted(p.name for p in pasta.glob("ppt/media/*")), midia, "mídia é compartilhada")
        r = rodar(NOVO_SLIDE, pasta, "--de", "slide3.xml")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(problemas_do_pacote(pasta), [])
        _, p = self.abrir(pasta)
        self.assertEqual(len(p.slides), 6)
        self.assertTrue(p.slides[4].shapes[-1].has_chart)
        self.assertEqual(p.slides[4].shapes[-1].chart.plots[0].categories[1], "B")

    def test_slide_em_branco_do_layout(self):
        pasta = self.preparar()
        from pptx import Presentation
        nome_layout = Presentation(self.tmp / "d.pptx").slide_layouts[1].name
        r = rodar(NOVO_SLIDE, pasta, "--layout", "slideLayout2")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("caixas reservadas", r.stdout)
        self.assertEqual(problemas_do_pacote(pasta), [])
        saida, p = self.abrir(pasta)
        novo = p.slides[-1]
        self.assertEqual(novo.slide_layout.name, nome_layout)
        self.assertGreaterEqual(len(novo.placeholders), 2)
        novo.shapes.title.text = "Escrito depois"
        p.save(saida)
        if TEM_SOFFICE:
            pdf = converter_pdf(saida, self.tmp)
            self.assertIsNotNone(pdf, "LibreOffice não converteu o deck com slide em branco")
            self.assertEqual(paginas_pdf(pdf), 5)

    def test_depois_poe_na_posicao(self):
        pasta = self.preparar()
        r = rodar(NOVO_SLIDE, pasta, "--de", "slide4", "--depois", "slide1")
        self.assertEqual(r.returncode, 0, r.stderr)
        _, p = self.abrir(pasta)
        self.assertEqual([s.shapes.title.text for s in p.slides],
                         ["Titulo 1", "Titulo 4", "Titulo 2", "Titulo 3", "Titulo 4"])

    def test_ids_unicos_em_varias_rodadas(self):
        pasta = self.preparar(n=2)
        vistos = set()
        for _ in range(3):
            r = rodar(NOVO_SLIDE, pasta, "--de", "slide1")
            self.assertEqual(r.returncode, 0, r.stderr)
            linha = [l for l in r.stdout.splitlines() if l.startswith("<p:sldId")][0]
            self.assertNotIn(linha, vistos)
            vistos.add(linha)
        self.assertEqual(problemas_do_pacote(pasta), [])
        _, p = self.abrir(pasta)
        self.assertEqual(len(p.slides), 5)

    def test_secao_recebe_o_slide_novo(self):
        pasta = self.preparar(n=2)
        arv = ler(pasta, "ppt/presentation.xml")
        raiz = arv.getroot()
        ext_lst = raiz.find(f"{{{NS_P}}}extLst")
        if ext_lst is None:
            ext_lst = etree.SubElement(raiz, f"{{{NS_P}}}extLst")
        ext = etree.SubElement(ext_lst, f"{{{NS_P}}}ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
        sec = etree.SubElement(etree.SubElement(ext, f"{{{NS_P14}}}sectionLst", nsmap={"p14": NS_P14}),
                               f"{{{NS_P14}}}section", name="Unica", id="{AAAAAAAA-2222-3333-4444-555555555555}")
        ids = etree.SubElement(sec, f"{{{NS_P14}}}sldIdLst")
        for s in raiz.iter(f"{{{NS_P}}}sldId"):
            etree.SubElement(ids, f"{{{NS_P14}}}sldId", id=s.get("id"))
        gravar(arv, pasta, "ppt/presentation.xml")
        r = rodar(NOVO_SLIDE, pasta, "--layout", "1")
        self.assertEqual(r.returncode, 0, r.stderr)
        raiz = ler(pasta, "ppt/presentation.xml").getroot()
        principais = [s.get("id") for s in raiz.iter(f"{{{NS_P}}}sldId")]
        da_secao = [s.get("id") for s in raiz.iter(f"{{{NS_P14}}}sldId")]
        self.assertEqual(principais, da_secao)

    def test_duplicar_e_limpar_volta_ao_original(self):
        pasta = self.preparar()
        antes = set(instantaneo(pasta))
        r = rodar(NOVO_SLIDE, pasta, "--de", "slide4")  # gráfico com planilha embutida
        self.assertEqual(r.returncode, 0, r.stderr)
        r = rodar(NOVO_SLIDE, pasta, "--de", "slide2")  # notas
        self.assertEqual(r.returncode, 0, r.stderr)
        tirar_da_lista(pasta, 6)
        tirar_da_lista(pasta, 5)
        r = rodar(LIMPAR, pasta)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(set(instantaneo(pasta)), antes)
        self.assertEqual(problemas_do_pacote(pasta), [])

    def test_erros_de_argumento(self):
        pasta = self.preparar(n=2)
        self.assertEqual(rodar(NOVO_SLIDE, pasta, "--de", "slide9").returncode, 3)
        self.assertEqual(rodar(NOVO_SLIDE, pasta, "--layout", "layout99").returncode, 3)
        self.assertEqual(rodar(NOVO_SLIDE, pasta, "--de", "1", "--layout", "1").returncode, 2)
        self.assertEqual(rodar(NOVO_SLIDE, pasta).returncode, 2)
        antes = instantaneo(pasta)
        self.assertEqual(rodar(NOVO_SLIDE, pasta, "--de", "1", "--depois", "slide9").returncode, 2)
        self.assertEqual(instantaneo(pasta), antes, "nada pode ser gravado quando --depois é inválido")
        self.assertEqual(rodar(NOVO_SLIDE, self.tmp / "nao_existe", "--de", "1").returncode, 3)


# ---------------------------------------------------------------- recalcular

def planilha_limpa(caminho):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vendas"
    for i, v in enumerate([10, 20, 30], 1):
        ws.cell(row=i, column=1, value=v)
        ws.cell(row=i, column=2, value=f"=A{i}*2")
    ws["A5"] = "=SUM(A1:A3)"
    ws["B5"] = "=AVERAGE(B1:B3)"
    wb.save(caminho)


def planilha_com_erros(caminho):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contas"
    ws["A1"] = 10
    ws["A2"] = 0
    ws["B1"] = "=A1/A2"          # #DIV/0!
    ws["B2"] = "=A1/0"           # #DIV/0!
    ws["B3"] = "=FUNCAOQUENAOEXISTE(A1)"  # #NAME?
    ws["B4"] = '="texto"+1'      # #VALUE!
    ws["B5"] = "=NA()"           # #N/A
    ws["B6"] = "=ACOS(2)"        # #NUM!
    ws["B7"] = "=A1+2"           # sem erro
    r = wb.create_sheet("Resumo")
    r["A1"] = "=Contas!A1*3"     # sem erro
    r["A2"] = "=#REF!+1"         # #REF!
    r["A3"] = "=SUM(Contas!A1:A2 Contas!B1:B2)"  # #NULL! (interseção vazia)
    wb.create_sheet("Vazia")
    wb.save(caminho)


def soffice_falso(pasta, corpo):
    pasta.mkdir(parents=True, exist_ok=True)
    s = pasta / "soffice"
    s.write_text("#!/bin/sh\n" + corpo + "\n")
    s.chmod(0o755)
    return pasta


@unittest.skipUnless(TEM_SOFFICE, "precisa do LibreOffice")
class TesteRecalcular(Base):
    def test_sem_erros_mantem_formulas_e_grava_valores(self):
        import openpyxl
        arq = self.tmp / "p.xlsx"
        planilha_limpa(arq)
        self.assertIsNone(openpyxl.load_workbook(arq, data_only=True)["Vendas"]["A5"].value)
        r = rodar(RECALCULAR, arq)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("5 fórmulas", r.stdout)
        valores = openpyxl.load_workbook(arq, data_only=True)["Vendas"]
        self.assertEqual(valores["A5"].value, 60)
        self.assertEqual(valores["B5"].value, 40)
        formulas = openpyxl.load_workbook(arq)["Vendas"]
        self.assertEqual(formulas["A5"].value.upper(), "=SUM(A1:A3)")

    def test_acha_todos_os_tipos_de_erro_em_json(self):
        arq = self.tmp / "e.xlsx"
        planilha_com_erros(arq)
        r = rodar(RECALCULAR, arq, "--json")
        self.assertEqual(r.returncode, 1, r.stdout + r.stderr)
        d = json.loads(r.stdout)
        self.assertTrue(d["recalculado"])
        self.assertEqual(d["situacao"], "com erros")
        abas = {a["aba"]: a for a in d["abas"]}
        self.assertEqual(set(abas), {"Contas", "Resumo", "Vazia"})
        contas = abas["Contas"]["erros"]
        self.assertEqual(contas["#DIV/0!"]["quantidade"], 2)
        self.assertEqual(contas["#DIV/0!"]["celulas"], ["B1", "B2"])
        for tipo, cel in (("#NAME?", "B3"), ("#VALUE!", "B4"), ("#N/A", "B5"), ("#NUM!", "B6")):
            self.assertEqual(contas[tipo]["celulas"], [cel], tipo)
        resumo = abas["Resumo"]["erros"]
        self.assertEqual(resumo["#REF!"]["celulas"], ["A2"])
        self.assertEqual(resumo["#NULL!"]["celulas"], ["A3"])
        self.assertEqual(abas["Vazia"]["total"], 0)
        self.assertEqual(d["total_erros"], 8)

    def test_saida_em_texto_por_aba(self):
        arq = self.tmp / "e.xlsx"
        planilha_com_erros(arq)
        r = rodar(RECALCULAR, arq)
        self.assertEqual(r.returncode, 1)
        self.assertIn('Aba "Contas": 6 erro(s)', r.stdout)
        self.assertIn("#DIV/0!: 2  B1, B2", r.stdout)
        self.assertIn('Aba "Vazia": sem erros', r.stdout)

    def test_mostrar_limita_a_lista(self):
        import openpyxl
        arq = self.tmp / "m.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        for i in range(1, 16):
            ws.cell(row=i, column=1, value=f"=1/0")
        wb.save(arq)
        d = json.loads(rodar(RECALCULAR, arq, "--json", "--mostrar", "5").stdout)
        info = d["abas"][0]["erros"]["#DIV/0!"]
        self.assertEqual(info["quantidade"], 15)
        self.assertEqual(len(info["celulas"]), 5)
        self.assertEqual(info["fora_da_lista"], 10)

    def test_vinculo_externo_recusa_sem_forcar(self):
        arq = self.tmp / "x.xlsx"
        planilha_limpa(arq)
        with zipfile.ZipFile(arq, "a") as z:
            z.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
        antes = arq.read_bytes()
        r = rodar(RECALCULAR, arq)
        self.assertEqual(r.returncode, 1)
        self.assertIn("--forcar", r.stderr)
        self.assertEqual(arq.read_bytes(), antes)

    def test_original_intacto_quando_soffice_falha(self):
        arq = self.tmp / "p.xlsx"
        planilha_limpa(arq)
        antes = arq.read_bytes()
        falso = soffice_falso(self.tmp / "falha", "exit 7")
        r = rodar(RECALCULAR, arq, env=dict(os.environ, PATH=f"{falso}:{os.environ['PATH']}"))
        self.assertEqual(r.returncode, 4)
        self.assertIn("original não foi mexido", r.stderr)
        self.assertEqual(arq.read_bytes(), antes)

    def test_limite_de_tempo(self):
        arq = self.tmp / "p.xlsx"
        planilha_limpa(arq)
        antes = arq.read_bytes()
        falso = soffice_falso(self.tmp / "lento", "sleep 20")
        r = rodar(RECALCULAR, arq, "--limite", "1", "--json",
                  env=dict(os.environ, PATH=f"{falso}:{os.environ['PATH']}"), timeout=15)
        self.assertEqual(r.returncode, 4)
        d = json.loads(r.stdout)
        self.assertFalse(d["recalculado"])
        self.assertIn("limite de 1 s", d["erro"])
        self.assertEqual(arq.read_bytes(), antes)

    def test_dependencias_e_argumentos(self):
        arq = self.tmp / "p.xlsx"
        planilha_limpa(arq)
        vazio = path_sem()
        self.assertEqual(rodar(RECALCULAR, arq, env=dict(os.environ, PATH=str(vazio))).returncode, 3)
        self.assertEqual(rodar(RECALCULAR, self.tmp / "nao.xlsx").returncode, 3)
        ruim = self.tmp / "ruim.xlsx"
        ruim.write_text("isto não é zip")
        self.assertEqual(rodar(RECALCULAR, ruim).returncode, 2)
        self.assertEqual(rodar(RECALCULAR, arq, "--limite", "0").returncode, 2)
        txt = self.tmp / "a.csv"
        txt.write_text("1,2")
        self.assertEqual(rodar(RECALCULAR, txt).returncode, 2)


class TesteAjuda(unittest.TestCase):
    def test_todos_tem_ajuda(self):
        for s in (MOSAICO, LIMPAR, NOVO_SLIDE, RECALCULAR):
            r = rodar(s, "-h")
            self.assertEqual(r.returncode, 0, s.name)
            self.assertIn("usage", r.stdout.lower())
            self.assertNotIn(chr(0x2014), s.read_text(encoding="utf-8"), f"travessão em {s.name}")


if __name__ == "__main__":
    unittest.main()
