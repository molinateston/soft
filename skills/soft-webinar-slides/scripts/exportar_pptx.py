#!/usr/bin/env python3
"""exportar_pptx.py: gera o deck.pptx EDITÁVEL a partir do deck.html, com as
notas do apresentador e a revelação por clique virando animação do PowerPoint.

Como funciona: abre o deck.html no Chromium (Playwright) com todos os cliques
abertos, lê do DOM cada bloco de cor, borda, imagem e texto (posição em pixels
do palco 1920x1080) e reconstrói como forma nativa (caixa de texto, retângulo,
imagem). O que entra no clique N vira um grupo com a animação do mesmo efeito:
  subir    -> Flutuar para dentro (presetID 42)
  aparecer -> Esmaecer (presetID 10)
  riscar   -> Revelar da esquerda (presetID 22)
  selo     -> Zoom (presetID 53)
Escala: 1920px = 13,333 polegadas, então 1px = 0,5pt (24px viram 12pt).
SVG e todo bloco marcado com data-imagem (desenho do agente) entram como imagem
na mesma posição e no mesmo clique: o texto em volta segue editável.

O texto fica editável; a fonte do perfil (fonte_pptx) precisa existir no
computador de quem abre. A fidelidade é aproximada (quebra de linha pode mudar
com outra fonte). Ver references/animacao-e-exportacao.md.

Uso:
  python3 scripts/exportar_pptx.py <pasta-ou-deck.html> [--saida deck.pptx]
"""
import argparse
import json
import re
import shutil
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU_PX = 914400 / 144  # 144 px por polegada

JS = r"""
(i) => {
  const s = document.querySelectorAll('.slide')[i];
  const R = s.getBoundingClientRect();
  const parse = (c) => { const m = (c||'').match(/rgba?\(([^)]+)\)/); if (!m) return null; const p = m[1].split(',').map(parseFloat); if (p.length > 3 && p[3] < 0.05) return null; return [p[0],p[1],p[2]].map(v => Math.round(v).toString(16).padStart(2,'0')).join('').toUpperCase(); };
  const clickOf = (el) => { const c = el.closest('[data-click]'); return c ? [+c.dataset.click, c.dataset.fx || 'subir'] : [0, null]; };
  const rotOf = (el) => { let rot = 0; for (let e = el; e && e !== s; e = e.parentElement) { const t = getComputedStyle(e).transform; const m = t && t.match(/matrix\(([^)]+)\)/); if (m) { const [a,b] = m[1].split(',').map(parseFloat); rot += Math.atan2(b,a) * 180 / Math.PI; } } return rot; };
  const geo = (el) => { const r = el.getBoundingClientRect(); const rot = rotOf(el);
    if (Math.abs(rot) < 0.5) return {x: r.x - R.x, y: r.y - R.y, w: r.width, h: r.height, rot: 0};
    const cx = r.x + r.width/2 - R.x, cy = r.y + r.height/2 - R.y; const w = el.offsetWidth, h = el.offsetHeight;
    return {x: cx - w/2, y: cy - h/2, w, h, rot}; };
  const items = []; let shots = 0;
  const vis = (el) => { const cs = getComputedStyle(el); return cs.display !== 'none' && cs.visibility !== 'hidden'; };
  // blocos de cor e borda
  s.querySelectorAll('*').forEach((el) => {
    if (el.closest('aside.notes') || !vis(el) || el.tagName === 'IMG') return;
    if (el.closest('svg, [data-imagem]')) {
      const raiz = el.closest('[data-imagem]') || el.closest('svg');
      if (raiz === el || (el.tagName.toLowerCase() === 'svg' && !el.parentElement.closest('svg, [data-imagem]'))) {
        const g = geo(el); const [click, fx] = clickOf(el);
        if (g.w >= 1 && g.h >= 1) { el.dataset.shot = 'x' + (shots++); items.push({k: 'shot', ...g, rot: 0, shot: el.dataset.shot, click, fx}); }
      }
      return;
    }
    const cs = getComputedStyle(el); const g = geo(el); if (g.w < 1 || g.h < 0.5) return;
    const fill = parse(cs.backgroundColor);
    const sides = ['Top','Right','Bottom','Left'].map(k => ({k, w: parseFloat(cs['border'+k+'Width'])||0, c: parse(cs['border'+k+'Color']), st: cs['border'+k+'Style']}));
    const all = sides.every(b => b.w > 0 && b.w === sides[0].w && b.c === sides[0].c);
    const [click, fx] = clickOf(el);
    const rad = parseFloat(cs.borderTopLeftRadius) || 0;
    if (fill || all) items.push({k: 'rect', ...g, fill, line: all ? sides[0].c : null, lineW: all ? sides[0].w : 0, dash: all && sides[0].st === 'dashed', rad, click, fx});
    if (!all) sides.forEach(b => { if (b.w > 0 && b.c && b.st !== 'none') {
      const r = {Top: {x: g.x, y: g.y, w: g.w, h: b.w}, Bottom: {x: g.x, y: g.y + g.h - b.w, w: g.w, h: b.w}, Left: {x: g.x, y: g.y, w: b.w, h: g.h}, Right: {x: g.x + g.w - b.w, y: g.y, w: b.w, h: g.h}}[b.k];
      items.push({k: 'rect', ...r, rot: 0, fill: b.c, line: null, lineW: 0, rad: 0, click, fx}); } });
  });
  // imagens
  s.querySelectorAll('img').forEach((im) => { if (!vis(im) || im.closest('[data-imagem]')) return; const g = geo(im); const [click, fx] = clickOf(im);
    items.push({k: 'img', ...g, src: im.getAttribute('src'), nw: im.naturalWidth, nh: im.naturalHeight, click, fx}); });
  // texto: um bloco por elemento não inline que tem texto direto ou via filhos inline
  const blocoDe = (el) => { for (let e = el; e && e !== s; e = e.parentElement) { if (getComputedStyle(e).display !== 'inline') return e; } return s; };
  const blocos = new Map();
  const walk = document.createTreeWalker(s, NodeFilter.SHOW_TEXT);
  let t;
  while ((t = walk.nextNode())) {
    if (!t.nodeValue.trim()) continue; const el = t.parentElement;
    if (el.closest('aside.notes') || el.closest('svg, [data-imagem]') || !vis(el)) continue;
    const b = blocoDe(el); if (!blocos.has(b)) blocos.set(b, []); blocos.get(b).push(t);
  }
  blocos.forEach((nodes, b) => {
    const rg = document.createRange(); rg.setStartBefore(nodes[0]); rg.setEndAfter(nodes[nodes.length-1]);
    const r = rg.getBoundingClientRect(); const cb = getComputedStyle(b);
    const runs = nodes.map((n, i) => { const cs = getComputedStyle(n.parentElement);
      let txt = n.nodeValue.replace(/\s+/g, ' ');
      if (i > 0) { const qb = document.createRange(); qb.setStartAfter(nodes[i-1]); qb.setEndBefore(n);
        if (qb.cloneContents().querySelector('br') && !/\s$/.test(nodes[i-1].nodeValue) && !/^\s/.test(txt)) txt = ' ' + txt; }  // <br> some no PPTX: vira espaço
      if (cs.textTransform === 'uppercase') txt = txt.toUpperCase();
      let deco = false; for (let e = n.parentElement; e && e !== s; e = e.parentElement) { if ((getComputedStyle(e).textDecorationLine||'').includes('line-through')) { deco = true; break; } }
      return {txt, color: parse(cs.color) || '000000', bold: (parseInt(cs.fontWeight,10)||400) >= 600, italic: cs.fontStyle === 'italic',
        size: parseFloat(cs.fontSize), font: cs.fontFamily.split(',')[0].replace(/['"]/g,'').trim(), spc: parseFloat(cs.letterSpacing)||0, strike: deco}; });
    runs[0].txt = runs[0].txt.replace(/^\s+/, ''); runs[runs.length-1].txt = runs[runs.length-1].txt.replace(/\s+$/, '');
    const [click, fx] = clickOf(b);
    const lh = Math.max(...nodes.map((n) => { const cs = getComputedStyle(n.parentElement); return parseFloat(cs.lineHeight) || parseFloat(cs.fontSize) * 1.15; }));
    const al = cb.textAlign === 'center' ? 'center' : (cb.textAlign === 'right' || cb.textAlign === 'end') ? 'right' : 'left';
    items.push({k: 'text', x: r.x - R.x, y: r.y - R.y, w: r.width, h: r.height, rot: rotOf(b), runs, lh, align: al, click, fx});
  });
  const notes = s.querySelector('aside.notes');
  return {bg: parse(getComputedStyle(s).backgroundColor) || 'FFFFFF', items, nota: notes ? notes.textContent.trim() : ''};
}
"""

P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def efeito(ids, spid, fx):
    """XML de um efeito de entrada por clique, alvo = forma (grupo) spid."""
    a, b, c = next(ids), next(ids), next(ids)
    vis = (f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
           f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
           f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')

    def anim(attr, de, para, dur=500):
        i = next(ids)
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{i}" dur="{dur}" fill="hold"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="{de}"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="{para}"/></p:val></p:tav></p:tavLst></p:anim>')

    def filtro(f, dur=500):
        i = next(ids)
        return (f'<p:animEffect transition="in" filter="{f}"><p:cBhvr><p:cTn id="{i}" dur="{dur}"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')

    if fx == "aparecer":
        pid, sub, corpo = 10, 0, vis + filtro("fade")
    elif fx == "riscar":
        pid, sub, corpo = 22, 8, vis + filtro("wipe(left)")
    elif fx == "selo":
        pid, sub, corpo = 53, 16, vis + anim("ppt_w", "0", "#ppt_w") + anim("ppt_h", "0", "#ppt_h") + filtro("fade")
    else:  # subir
        pid, sub, corpo = 42, 0, vis + filtro("fade") + anim("ppt_x", "#ppt_x", "#ppt_x") + anim("ppt_y", "#ppt_y+.1", "#ppt_y")
    o1, o2 = next(ids), next(ids)
    return (f'<p:par><p:cTn id="{o1}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{o2}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{a}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold" nodeType="clickEffect">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{corpo}</p:childTnLst></p:cTn></p:par>'
            f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


def timing(cliques):
    def gen():
        n = 3
        while True:
            yield n
            n += 1
    ids = gen()
    pars = "".join(efeito(ids, spid, fx) for spid, fx in cliques)
    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{pars}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
           f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    return etree.fromstring(xml)


def px(v):
    return Emu(int(round(v * EMU_PX)))


def rgb(h):
    return RGBColor.from_string(h)


def montar(dados, perfil, base_dir, saida):
    fonte_html = perfil["fonte"].split(",")[0].strip().strip("'\"")
    acao_html = perfil["fonte_acao"].split(",")[0].strip().strip("'\"")
    mapa = {fonte_html: perfil.get("fonte_pptx", fonte_html), acao_html: perfil.get("fonte_acao_pptx", acao_html),
            "Liberation Sans": perfil.get("fonte_pptx", "Arial"), "DejaVu Serif": perfil.get("fonte_acao_pptx", "Georgia")}
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1920), px(1080)
    blank = prs.slide_layouts[6]
    resumo = []
    for d in dados:
        sl = prs.slides.add_slide(blank)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(d["bg"])
        grupos = {}
        for it in d["items"]:
            k = it["click"]
            if k and k not in grupos:
                grupos[k] = (sl.shapes.add_group_shape(), it["fx"])
            alvo = grupos[k][0].shapes if k else sl.shapes
            if it["k"] == "rect":
                r = it["rad"]
                m = min(it["w"], it["h"])
                forma = MSO_SHAPE.ROUNDED_RECTANGLE if r > 0.5 else MSO_SHAPE.RECTANGLE
                shp = alvo.add_shape(forma, px(it["x"]), px(it["y"]), px(it["w"]), px(max(it["h"], 1)))
                if forma == MSO_SHAPE.ROUNDED_RECTANGLE and m > 0:
                    shp.adjustments[0] = min(0.5, r / m)
                if it["fill"]:
                    shp.fill.solid()
                    shp.fill.fore_color.rgb = rgb(it["fill"])
                else:
                    shp.fill.background()
                if it["line"]:
                    shp.line.color.rgb = rgb(it["line"])
                    shp.line.width = Pt(it["lineW"] * 0.5)
                    if it.get("dash"):
                        shp.line.dash_style = MSO_LINE.DASH
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
                estilo = shp._element.find(qn("p:style"))
                if estilo is not None:  # o estilo padrão do tema põe sombra e contorno que o deck não tem
                    shp._element.remove(estilo)
                if it["rot"]:
                    shp.rotation = it["rot"]
                shp.text_frame.text = ""
            elif it["k"] == "shot":  # SVG e desenho marcado com data-imagem: entram como imagem, na mesma posição
                if it.get("png") and Path(it["png"]).exists():
                    alvo.add_picture(it["png"], px(it["x"]), px(it["y"]), px(it["w"]), px(it["h"]))
            elif it["k"] == "img":
                src = (base_dir / it["src"]).resolve()
                if not src.exists():
                    continue
                w, h = it["w"], it["h"]
                if it["nw"] and it["nh"]:
                    esc = min(w / it["nw"], h / it["nh"])
                    w2, h2 = it["nw"] * esc, it["nh"] * esc
                    alvo.add_picture(str(src), px(it["x"] + (w - w2) / 2), px(it["y"] + (h - h2) / 2), px(w2), px(h2))
                else:
                    alvo.add_picture(str(src), px(it["x"]), px(it["y"]), px(w), px(h))
            else:
                folga = 0.04 * it["w"] + 6
                x = it["x"] - (folga / 2 if it["align"] == "center" else (folga if it["align"] == "right" else 0))
                tb = alvo.add_textbox(px(x), px(it["y"]), px(it["w"] + folga), px(it["h"] + 4))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.TOP
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                p = tf.paragraphs[0]
                p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(it["align"], PP_ALIGN.LEFT)
                p.line_spacing = Pt(it["lh"] * 0.5)
                for rn in it["runs"]:
                    if not rn["txt"]:
                        continue
                    r = p.add_run()
                    r.text = rn["txt"]
                    f = r.font
                    f.size = Pt(rn["size"] * 0.5)
                    f.bold = rn["bold"]
                    f.italic = rn["italic"]
                    f.name = mapa.get(rn["font"], rn["font"])
                    f.color.rgb = rgb(rn["color"])
                    rpr = r._r.get_or_add_rPr()
                    if rn["spc"]:
                        rpr.set("spc", str(int(rn["spc"] * 0.5 * 100)))
                    if rn["strike"]:
                        rpr.set("strike", "sngStrike")
                if it["rot"]:
                    tb.rotation = it["rot"]
        if grupos:
            ordem = sorted(grupos)
            arvore = sl.shapes._spTree
            for k in ordem:  # o que entra por clique fica por cima da base, na ordem dos cliques
                el = grupos[k][0]._element
                arvore.remove(el)
                arvore.append(el)
            sl._element.append(timing([(grupos[k][0].shape_id, grupos[k][1]) for k in ordem]))
        sl.notes_slide.notes_text_frame.text = d["nota"]
        resumo.append(len(grupos))
    prs.save(saida)
    return resumo


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("alvo")
    ap.add_argument("--saida")
    a = ap.parse_args()
    alvo = Path(a.alvo)
    html = alvo / "deck.html" if alvo.is_dir() else alvo
    perfil_p = html.parent / ".perfil-usado.json"
    perfil = json.loads(perfil_p.read_text(encoding="utf-8")) if perfil_p.exists() else json.loads(
        (Path(__file__).resolve().parent.parent / "assets" / "perfil-padrao.json").read_text(encoding="utf-8"))
    saida = Path(a.saida) if a.saida else html.parent / "deck.pptx"
    from playwright.sync_api import sync_playwright
    dados = []
    with sync_playwright() as p:
        try:
            br = p.chromium.launch()
        except Exception:
            br = p.chromium.launch(executable_path="/snap/bin/chromium")
        pg = br.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=1)
        pg.goto(html.resolve().as_uri())
        pg.wait_for_function("window.deckPronto === true", timeout=20000)
        pg.add_style_tag(content="*{transition:none!important;animation:none!important}#painel-notas{display:none!important}")
        n = pg.evaluate("window.deck.count")
        tmp = Path(tempfile.mkdtemp(prefix="pptx-desenho-"))
        for i in range(n):
            pg.evaluate(f"window.deck.show({i})")
            pg.wait_for_timeout(50)
            d = pg.evaluate(JS, i)
            for it in d["items"]:
                if it["k"] == "shot":
                    it["png"] = str(tmp / f"s{i}-{it['shot']}.png")
                    try:
                        pg.locator(f'.slide.active [data-shot="{it["shot"]}"]').screenshot(path=it["png"], omit_background=True)
                    except Exception:
                        it["png"] = ""
            dados.append(d)
        br.close()
    resumo = montar(dados, perfil, html.parent, str(saida))
    shutil.rmtree(tmp, ignore_errors=True)
    print(f"{len(dados)} slides -> {saida} (cliques animados por slide: {resumo})")


if __name__ == "__main__":
    main()
